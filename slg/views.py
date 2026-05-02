import requests
import csv
import wikipedia
import wikipediaapi
import operator
# import pandas as pd
from typing import Iterable
from django.contrib import messages 
from django.contrib.admin.models import ADDITION, CHANGE, LogEntry
from django.contrib.admin.views.main import TO_FIELD_VAR, IGNORED_PARAMS
from django.contrib.admin.filters import AllValuesFieldListFilter
from django.contrib.auth.decorators import login_required
from django.core import serializers
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.db import connection, transaction, IntegrityError
from django.db.models import Prefetch, Q, Count, Max, Min, Avg, F, Exists, OuterRef, Subquery, Case, When, Value, IntegerField
from django.http import HttpResponse, JsonResponse, HttpResponseRedirect, StreamingHttpResponse
from django.forms import inlineformset_factory, modelformset_factory
from django.shortcuts import render , get_object_or_404, redirect
from django.urls import reverse_lazy, reverse
from django.views.generic import UpdateView
from .models import *
from .admin import process_download_infos
from .serializers import *
from .forms import *
from .resources import ObjLSNOResource
# from .models import Slg, Obj, Obj_Person, Obj_Ref, PersonFunktion, Person, Mzstaette, Metall, AvBildtyp, RvBildtyp, RvBildtyp_Schlagwort, AvBildtyp_Schlagwort, Schlagwort, Muenztyp, Mztyp_Person, Faelschung
from rest_framework import viewsets, generics
from rest_framework.authentication import TokenAuthentication, SessionAuthentication
from rest_framework.decorators import api_view, permission_classes, authentication_classes
from rest_framework.permissions import IsAuthenticated, IsAuthenticatedOrReadOnly   # nur eingeloggte
from rest_framework.response import Response
from bootstrap_modal_forms.generic import BSModalCreateView, BSModalUpdateView

from urllib.parse import quote_plus, unquote, unquote_plus
from rdflib import Graph, Namespace, Literal
from rdflib.namespace import RDF, FOAF, DCTERMS, URIRef

from collections import Counter, OrderedDict, defaultdict
from itertools import chain

from functools import reduce
from operator import or_, and_

from .filter_config import FILTER_PARAMETERS

from django.db.models import Q
from django.contrib.contenttypes.models import ContentType
from django.utils import timezone
from .authentication import QueryParamTokenAuthentication
import traceback
import io
import json  # Füge diesen Import hinzu
import logging
import re  # Für Regular Expressions
import shlex  # Für Anführungszeichen-behaftete Strings
from django.core.cache import cache
from django.template.loader import render_to_string
from rest_framework.views import APIView
from rest_framework.permissions import AllowAny
from .helpers import *
from .filters import apply_filters, build_filters, is_valid_qparam
from .helpers import get_matrix_from_muenztypen
from django.http import JsonResponse
from django.shortcuts import render
from .helpers import get_timeline_data

logger = logging.getLogger(__name__)

def replace_monogram_references(text):
    """
    Ersetzt Monogramm-Referenzen in eckigen Klammern durch entsprechende SVG-Symbole.
    
    Beispiel: "[Lorber Monogram 219]" wird durch das SVG-Symbol für "Lorber Monogram 219" ersetzt.
    """
    if not text:
        return text
    
    # Finde alle [Name] Muster im Text
    pattern = r'\[([^\]]+)\]'
    matches = re.findall(pattern, text)
    
    for match in matches:
        try:
            # Suche nach dem OffizinSymbol mit dem Namen
            symbol = OffizinSymbol.objects.get(name=match)
            # Ersetze [Name] durch das SVG-HTML
            text = text.replace(f'[{match}]', symbol.get_svg_html())
        except OffizinSymbol.DoesNotExist:
            # Wenn das Symbol nicht gefunden wird, lasse den Text unverändert
            # Optional: Logge eine Warnung
            logger.warning(f"OffizinSymbol mit dem Namen '{match}' nicht gefunden.")
            continue
        except Exception as e:
            # Für andere Fehler, logge sie aber lasse den Text unverändert
            logger.error(f"Fehler beim Ersetzen von '{match}': {e}")
            continue
    
    return text

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen
def get_adjacent_invnrs(request):
    current_invnr = request.GET.get('invnr')
    slg_id = request.GET.get('slg_id')
    slgteil_id = request.GET.get('slgteil_id')
    
    # Log the received parameters for debugging
    print(f"Parameters received: invnr='{current_invnr}', slg_id='{slg_id}', slgteil_id='{slgteil_id}'")
    
    if not current_invnr or not slg_id:
        return Response({'error': 'Inventarnummer und Sammlungs-ID sind erforderlich'}, status=400)
    
    try:
        # Ensure we're using valid IDs for slg and slgteil
        slg_id_int = int(slg_id)
        slgteil_id_int = int(slgteil_id) if slgteil_id else None
        
        # Query logic
        filter_kwargs = {'Slg_id': slg_id_int}
        if slgteil_id_int:
            filter_kwargs['SlgTeil_id'] = slgteil_id_int
            
        objects = Obj.objects.filter(**filter_kwargs).order_by('invnr')
        
        # Get all inventory numbers as they are (as strings)
        invnr_list = list(objects.values_list('invnr', flat=True))
        
        # Debug information
        print(f"Found {len(invnr_list)} objects with these filters")
        print(f"First 5 inventory numbers: {invnr_list[:5] if len(invnr_list) >= 5 else invnr_list}")
        print(f"Looking for inventory number: '{current_invnr}'")
        
        # Check if invnr exists in the list
        if current_invnr not in invnr_list:
            return Response({
                'error': f"Inventarnummer '{current_invnr}' nicht gefunden in der gefilterten Liste.",
                'available_invnrs': invnr_list[:10]  # Show first 10 as hint
            }, status=404)
        
        # Find the index while keeping invnr as string
        current_index = invnr_list.index(current_invnr)
        print(f"Found at index {current_index}")
        
        # Get object IDs that match these inventory numbers
        obj_map = {obj.invnr: obj.id for obj in objects}
        
        # Return result with proper bounds checking
        result = {
            'prev2': invnr_list[current_index - 2] if current_index >= 2 else None,
            'prev2_id': obj_map.get(invnr_list[current_index - 2]) if current_index >= 2 else None,
            'prev1': invnr_list[current_index - 1] if current_index >= 1 else None,
            'prev1_id': obj_map.get(invnr_list[current_index - 1]) if current_index >= 1 else None,
            'current': current_invnr,
            'current_id': obj_map.get(current_invnr),
            'next1': invnr_list[current_index + 1] if current_index < len(invnr_list) - 1 else None,
            'next1_id': obj_map.get(invnr_list[current_index + 1]) if current_index < len(invnr_list) - 1 else None,
            'next2': invnr_list[current_index + 2] if current_index < len(invnr_list) - 2 else None,
            'next2_id': obj_map.get(invnr_list[current_index + 2]) if current_index < len(invnr_list) - 2 else None,
        }
        
        return Response(result)
        
    except ValueError as e:
        # More detailed error message for debugging
        return Response({
            'error': f"ValueError: {str(e)}. Parameter types: invnr={type(current_invnr).__name__}, "
                     f"slg_id={type(slg_id).__name__}, slgteil_id={type(slgteil_id).__name__}"
        }, status=400)
    
    except Exception as e:
        # Catch any other errors
        return Response({
            'error': f"Unexpected error: {str(e)}"
        }, status=500)

class SlgTeilListAPIView(generics.ListAPIView):
    serializer_class = SlgTeilSerializer

    def get_queryset(self):
        slg_id = self.kwargs.get('slg_id')
        return SlgTeil.objects.filter(
            Q(obj__Slg__id=slg_id) |
            Q(idfk_Slg_SlgTeil__id=slg_id)
        ).distinct()

@api_view(['GET'])
@authentication_classes([QueryParamTokenAuthentication, SessionAuthentication, TokenAuthentication])
@permission_classes([IsAuthenticated])
def export_xlsx(request):
    filters = {key: value for key, value in request.GET.items() if key not in IGNORED_PARAMS}

    queryset = Obj.objects.all()
    invalid_filter_fields = []
    for key, value in filters.items():
        # Split the key to extract the field name and lookup type (if any)
        field_lookup_parts = key.split('__')
        field_name = field_lookup_parts[0]
        lookup_type = field_lookup_parts[1] if len(field_lookup_parts) > 1 else 'exact'

        # Unknown fields should not crash the export endpoint
        try:
            model_field = Obj._meta.get_field(field_name)
        except Exception:
            invalid_filter_fields.append(field_name)
            continue

        # Handle '__exact' lookup for ForeignKey fields correctly
        if lookup_type == 'exact' and isinstance(model_field, models.ForeignKey):
            filter_key = f"{field_name}_id"
        else:
            filter_key = key

        # Convert the value if it's a boolean field or if it's an '__isnull' lookup
        if lookup_type == 'isnull':
            filter_value = value == 'True'
        elif model_field.get_internal_type() == 'BooleanField':
            filter_value = value == 'on'
        else:
            filter_value = value

        # Apply the filter
        if filter_value is not None:
            try:
                queryset = queryset.filter(**{filter_key: filter_value})
            except Exception:
                # Ignore invalid lookup/value combinations instead of 500
                continue

    if invalid_filter_fields:
        invalid_filter_fields = sorted(set(invalid_filter_fields))
        return JsonResponse({"error": "Ungültige Filterfelder", "fields": invalid_filter_fields}, status=400)

    resource = ObjLSNOResource()
    dataset = resource.export(queryset)
    response = HttpResponse(dataset.xlsx, content_type='application/vnd.ms-excel')
    filename_base = 'objekte' if not filters else 'obj_filter'
    response['Content-Disposition'] = f'attachment; filename="{filename_base}_export.xlsx"'
    return response

class ObjCreateView(BSModalCreateView):
    template_name = 'slg/create_object.html'
    form_class = ObjForm
    success_message = 'Success: Objekt was created.'
    success_url = reverse_lazy('Objektliste')

# Update

class ObjUpdateView(BSModalUpdateView):
    model = Obj
    template_name = 'slg/update_obj.html'
    form_class = ObjForm
    success_message = 'Success: Obj was updated.'
    success_url = reverse_lazy('Objektliste')

@login_required
def MzUpdate(request, id):
   obj = Obj.objects.get(id=id)
   PersonFormSet = inlineformset_factory(Obj, Obj_Person, fields=('idfk_Person','idfk_PersonFunktion','appears_on_rev'), can_delete=True, extra=2)
   RefFormSet = inlineformset_factory(Obj, Obj_Ref, fields=('idfk_Ref','nummer', 'variante', 'link'), can_delete=True, extra=1,)
   # RefFormSet = inlineformset_factory(Obj, Obj_Ref, form=RefUpdateForm, can_delete=True, extra=1,)
   #print(formset)
   form = DetailUpdateForm(instance=obj)
   if request.method == 'POST':
      #print("erstes IF")
      #print(request.POST)
      Personen = PersonFormSet(request.POST, instance=obj, prefix='personen')
      Literatur = RefFormSet(request.POST, instance=obj, prefix='referenzen')
      form = DetailUpdateForm(request.POST, instance=obj)
      if form.is_valid() and Literatur.is_valid() and Personen.is_valid():
         #print("zweites IF")
         form.save()
         Literatur.save()
         Personen.save()
         # Personen = Personen.save(commit=False)
         # for inline_form in Personen:
         #        #if inline_form.cleaned_data:
         #            #person = inline_form.save(commit=False)
         #            inline_form.save()
         #formset.save()
         # Literatur = Literatur.save(commit=False)
         # for inline_form in Literatur:
         #        #if inline_form.cleaned_data:
         #            #person = inline_form.save(commit=False)
         #            inline_form.save()
         #formset.save()
         return HttpResponseRedirect(request.path_info)
   else:
      Personen = PersonFormSet(instance=obj, prefix='personen')
      Literatur = RefFormSet(instance=obj, prefix='referenzen')
   avtype = AvBildtyp.objects.all().order_by('name')
   rvtype = RvBildtyp.objects.all().order_by('name')
   
   
   
   
   context = {
      'form': form,
      'Personen': Personen,
      'Refs': Literatur,
      'test': obj,
      'Avtypes': avtype,
      'Rvtypes': rvtype,
      
   }

   
      # if form.is_valid():
      #    form.save()
         #return redirect('/')
   # def form_valid(self, form):
   #    print(form.cleaned_data)
   #    return super().form_valid(form)

   # def get_object(self):
   #    id_ = self.kwargs.get("id")
   #    return get_object_or_404(Obj, id=id_)

   return render(request, 'slg/obj_form.html', context)

def index(request):
    latest_entries = MuenztypObjektAnzeige.objects.exclude(obj_id__isnull=True).order_by('-last_modified')[:10]

    browse_url = reverse('Objektliste')

    def build_wordcloud_data(default_field, undetermined_field, filter_name, limit=80):
        counts = Counter()

        for keyword in chain(
            Obj.objects.filter(Typ__isnull=False).values_list(default_field, flat=True).iterator(),
            Obj.objects.filter(Typ__isnull=True).values_list(undetermined_field, flat=True).iterator(),
        ):
            if not keyword:
                continue
            normalized_keyword = str(keyword).strip()
            if not normalized_keyword:
                continue
            counts[normalized_keyword] += 1

        return [
            {
                'text': keyword,
                'value': count,
                'url': f"{browse_url}?{filter_name}={quote_plus(keyword)}",
            }
            for keyword, count in sorted(counts.items(), key=lambda item: (-item[1], item[0].lower()))[:limit]
        ]

    av_wordcloud_data = build_wordcloud_data(
        'Typ__av_bildtyp__avbildtyp_schlagwort__schlagwort__name',
        'av_bildtyp__avbildtyp_schlagwort__schlagwort__name',
        'av_schlagwort',
    )
    rv_wordcloud_data = build_wordcloud_data(
        'Typ__rv_bildtyp__rvbildtyp_schlagwort__schlagwort__name',
        'rv_bildtyp__rvbildtyp_schlagwort__schlagwort__name',
        'rv_schlagwort',
    )

    context = {
        'latest_entries': latest_entries,
        'av_wordcloud_data': av_wordcloud_data,
        'rv_wordcloud_data': rv_wordcloud_data,
    }
    
    return render(request, 'slg/index.html', context)

def neuerschliessungen(request):
    latest_entries_qs = (
        MuenztypObjektAnzeige.objects
        .exclude(obj_id__isnull=True)
        .order_by('-last_modified')
    )

    page = request.GET.get('page', 1)
    paginator = Paginator(latest_entries_qs, 24)

    try:
        latest_entries = paginator.page(page)
    except PageNotAnInteger:
        latest_entries = paginator.page(1)
    except EmptyPage:
        latest_entries = paginator.page(paginator.num_pages)

    return render(request, 'slg/neuerschliessungen.html', {
        'latest_entries': latest_entries,
    })

def sammlungen_uebersicht(request):
    sammlungen = (
        Slg.objects
        .only('id', 'name', 'beschreibung', 'cover')
        .annotate(obj_count=Count('obj'))
        .order_by('-obj_count', 'name')
    )
    return render(request, 'slg/sammlungen_uebersicht.html', {
        'sammlungen': sammlungen,
    })

# Neue API-Endpoints für die ausgelagerten Daten

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen
def statistics_api(request):
    """API-Endpoint für Sammlungsstatistiken."""
    anzahl_qs = Obj.objects.count()
    max_min_datierung = Muenztyp.objects.aggregate(Max('dat_von'), Min('dat_von'))
    max_min_gewicht = Obj.objects.aggregate(
        Max('gewicht'), Min('gewicht'),
        Max('durchmesser'), Min('durchmesser')
    )
    anzahl_mt = Muenztyp.objects.count()
    anzahl_person = Person.objects.count()
    anzahl_mstaetten = Mzstaette.objects.count()
    anzahl_kontrolliert = Obj.objects.filter(workflow__name="kontrolliert").count()
    anzahl_zurkontrolle = Obj.objects.filter(workflow__name="Bereit zur Kontrolle").count()
    
    return Response({
        'anzahl_qs': anzahl_qs,
        'anzahl_mt': anzahl_mt,
        'anzahl_person': anzahl_person,
        'anzahl_mstaetten': anzahl_mstaetten,
        'max_min_datierung': max_min_datierung,
        'max_min_gewicht': max_min_gewicht,
        'anzahl_kontrolliert': anzahl_kontrolliert,
        'anzahl_zurkontrolle': anzahl_zurkontrolle,
    })

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen
def sammlungen_api(request):
    """API-Endpoint für alle Sammlungen."""
    try:
        # Debug: Prüfe Anzahl der Sammlungen
        sammlungen_count = Slg.objects.count()
        logger.info(f"Gefundene Sammlungen: {sammlungen_count}")
        
        if sammlungen_count == 0:
            # Wenn keine Sammlungen vorhanden sind, geben wir ein Test-Objekt zurück
            logger.warning("Keine Sammlungen in der Datenbank!")
            return Response([{
                "id": 1,
                "name": "Test-Sammlung",
                "beschreibung": "Dies ist ein Test",
                "cover": "",
                "kategorie": {"id": 1, "name": "Test-Kategorie"}
            }])
        
        sammlungen = Slg.objects.all().select_related('kategorie')
        serializer = SlgSerializer(sammlungen, many=True)
        logger.info(f"Serialisierte Sammlungen: {len(serializer.data)}")
        return Response(serializer.data)
        
    except Exception as e:
        logger.error(f"Fehler in sammlungen_api: {str(e)}", exc_info=True)
        return Response(
            {"error": f"Fehler beim Laden der Sammlungen: {str(e)}"},
            status=500
        )

def NomismaInfo(link):
   url2 = "http://nomisma.org/query?query=PREFIX+skos%3A+%3Chttp%3A%2F%2Fwww.w3.org%2F2004%2F02%2Fskos%2Fcore%23%3E%0D%0A%0D%0ASELECT+*+WHERE+%7B%0D%0A++OPTIONAL+%7B%3C"+link+"%3E+skos%3AexactMatch+%3Flink%7D%0D%0A++%3C"+link+"%3E+skos%3Adefinition+%3Fdef.%0D%0A++FILTER+langMatches+%28lang%28%3Fdef%29%2C+%27en%27%29%0D%0A%7D&output=csv"
   print(url2)
   with requests.Session() as s:
      download = s.get(url2)
      decoded_content = download.content.decode('utf-8')
      exactmatch = csv.DictReader(decoded_content.splitlines(), delimiter=',')
      try:
         row2 = next(exactmatch)
         definition = row2['def']
      except:
         definition = ""
      return exactmatch, definition

def PersonModalView(request, id):
   person = get_object_or_404(Person, id=id)
   
   if person.WikiDe:
      link = unquote_plus(person.WikiDe)
      print(link)
      # encode(link, link)
      # print(link.decode('utf8').title())
      link = link.replace('https://de.wikipedia.org/wiki/', '')
      # wikipedia.set_lang("de")
      # wiki = wikipedia.summary(link)
      wiki_wiki = wikipediaapi.Wikipedia('de')
      page_py = wiki_wiki.page(link)
      wiki = page_py.summary
      #print(page_py.__dict__)
      # print ("Nr. of images on page: %s" % len(page_py.image))
      S = requests.Session()

      URL = "https://de.wikipedia.org/w/api.php"

      PARAMS = {
         "action": "query",
         "format": "json",
         "titles": link,
         "prop": "extracts",
         "prop": "pageimages",
         "pithumbsize": '300',
      }

      R = S.get(url=URL, params=PARAMS)
      DATA = R.json()

      PAGES = DATA['query']['pages']
      print(PAGES)

      for k, v in PAGES.items():
         try:
            thumbnail = str(v['thumbnail']['source'])
         except:
            thumbnail = ""
         
   elif person.WikiEn:
      link = unquote_plus(person.WikiEn)
      print(link)
      # encode(link, link)
      # print(link.decode('utf8').title())
      link = link.replace('https://en.wikipedia.org/wiki/', '')
      # wikipedia.set_lang("de")
      # wiki = wikipedia.summary(link)
      wiki_wiki = wikipediaapi.Wikipedia('en')
      page_py = wiki_wiki.page(link)
      wiki = page_py.summary
      S = requests.Session()

      URL = "https://en.wikipedia.org/w/api.php"

      PARAMS = {
         "action": "query",
         "format": "json",
         "titles": link,
         "prop": "extracts",
         "prop": "pageimages",
         "pithumbsize": '300',
      }

      R = S.get(url=URL, params=PARAMS)
      DATA = R.json()

      PAGES = DATA['query']['pages']
      print(PAGES)

      for k, v in PAGES.items():
         try:
            thumbnail = str(v['thumbnail']['source'])
         except:
            thumbnail = ""
   else:
      wiki = ""
      thumbnail = ""

   if person.name_nom_id:
      exactmatch, definition = NomismaInfo(person.name_nom_id)
   else:
      exactmatch = ""
      definition = ""
   # if person.VIAF:
   #    from xml.dom import minidom
   #    link = person.VIAF + "rdf.xml"
   #    with requests.Session() as s:
   #       download = s.get(link)
   #       decoded_content = download.content.decode('utf-8')
   #    xmldoc = minidom.parse(decoded_content)
      
   context = {
      'Person': person,
      'Wiki': wiki,
      'WikiBild': thumbnail,
      'Exactmatch': exactmatch,
      'Definition': definition,
   }
   
   return render(request, 'slg/modal_person.html', context)

def NominalModalView(request, id):
   nominal = get_object_or_404(Nominal, id=id)
   # if nominal.WikiDe:
   #    link = unquote_plus(nominal.WikiDe)
      
   #    # encode(link, link)
   #    # print(link.decode('utf8').title())
   #    link = link.replace('https://de.wikipedia.org/wiki/', '')
   #    wikipedia.set_lang("de")
   #    wiki = wikipedia.summary(link)
   # else:
   #    wiki = ""

   if nominal.name_nom_id:
      # url2 = "http://nomisma.org/query?query=PREFIX+skos%3A+%3Chttp%3A%2F%2Fwww.w3.org%2F2004%2F02%2Fskos%2Fcore%23%3E%0D%0A%0D%0ASELECT+*+WHERE+%7B%0D%0A++%3C"+nominal.name_nom_id+"%3E+skos%3AcloseMatch+%3Fo%0D%0A%7D&output=csv"
      exactmatch, definition = NomismaInfo(nominal.name_nom_id)
         # for row in cr:
         #    print(row['link'])
         # definition = cr
         # definition = list(definition)
         # print(definition[1][1])
         # definition = ""
         # definition = closematch[1][1]
         
         # print(cr['link'])
   context = {
      'Nominal': nominal,
      'Exactmatch': exactmatch,
      'Definition': definition,
   }
   
   return render(request, 'slg/modal_nominal.html', context)

def TypView(request, id):
    """
    Detailansicht für einen Münztyp.
    Holt den Typ selbst, alle für die linke Infospalte benötigten FK‑Felder
    und alle zugehörigen Objekte (inklusive deren wichtigsten FK‑Felder)
    in möglichst wenigen Datenbank‑Queries.
    """
 
    # ---------------------------------------------------------------
    # 1. Münztyp mit allen FK/M2M in wenigen Queries laden
    obj = get_object_or_404(
        Muenztyp.objects.select_related(
            'Nominal',
            'Mzstaette',
            'Muenzstand',
            'Reichskreis',
            'region',
            'Metall',
            'Objekttyp',
            'workflow',
            'av_bildtyp',
            'rv_bildtyp',
        ).prefetch_related(
            'mztyp_person_set__idfk_Person',
            'mztyp_person_set__idfk_PersonFunktion',
            'typ_ref_set__Ref',
            'Konkordanz__Ref',
            'av_bildtyp__avbildtyp_schlagwort_set__schlagwort',
            'rv_bildtyp__rvbildtyp_schlagwort_set__schlagwort',
        ),
        id=id,
    )
 
    # ---------------------------------------------------------------
    # 2. Bereits geprefetchte Daten abgreifen
    # ---------------------------------------------------------------
    personen = obj.mztyp_person_set.all()

    av_schlagworte = (
        obj.av_bildtyp.avbildtyp_schlagwort_set.all()
        if obj.av_bildtyp else []
    )
    rv_schlagworte = (
        obj.rv_bildtyp.rvbildtyp_schlagwort_set.all()
        if obj.rv_bildtyp else []
    )

    # ---------------------------------------------------------------
    # 3. Objekte über MuenztypObjektAnzeige laden (denormalisiert,
    #    keine JOINs nötig) mit DB-Level Pagination (LIMIT/OFFSET)
    # ---------------------------------------------------------------
    objekte_qs = (
        MuenztypObjektAnzeige.objects
            .filter(typ_fk_id=id)
            .only(
                "id", "obj_id", "Slg", "objekttitel",
                "datierung_verbale",
                "thumbnail_av_url", "thumbnail_rv_url",
            )
            .order_by("obj_id")
    )
    paginator = Paginator(objekte_qs, 12)
    page = request.GET.get("page", 1)
    try:
        objekte_page = paginator.page(page)
    except PageNotAnInteger:
        objekte_page = paginator.page(1)
    except EmptyPage:
        objekte_page = paginator.page(paginator.num_pages)

    # ---------------------------------------------------------------
    # 4. Template Kontext
    # ---------------------------------------------------------------
    context = {
        "Muenztyp": obj,
        "objekte_page": objekte_page,
        "MztypPersonen": personen,
        "avschlagworte": av_schlagworte,
        "rvschlagworte": rv_schlagworte,
    }

    return render(request, "slg/typ_detail.html", context)

def SammlungView(request, id):
   
   sammlung = get_object_or_404(Slg, id=id)
   # avbildtypen = AvBildtyp.objects.prefetch_related('avmztypen','avmztypen__objekte','avmztypen__objekte__SlgTeil').filter(schlagworte=id,)
   # typen = Muenztyp.objects.select_related('rv_bildtyp_set',).filter(rv_bildtyp__schlagworte=id,)
   # rvbildtypen = typen.values_list('rv_bildtyp__name', 'rv_bildtyp__id',).exclude(rv_bildtyp__name=None).distinct()
   # print(rvbildtypen.muenztyp.all())
   
   context = {
      'Sammlung': sammlung,
   }
   
   return render(request, 'slg/Sammlung.html', context)

def AvSchlagwortView(request, id):
   
   # obj = Muenztyp.objects.select_related('Muenzstand', 'Herstellung', 'Nominal', 'Mzstaette', 'Metall', 'region', 'av_bildtyp', 'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen', 'Objekttyp',).prefetch_related('obj_set', 'mztyp_person_set', 'typ_ref_set',).get(id=id)
   # rvbildtypen = RvBildtyp.objects.prefetch_related('muenztyp_set','obj_set').exclude(name__isnull=True).filter(schlagworte=id,)
   avbildtypen = AvBildtyp.objects.prefetch_related('avmztypen','avmztypen__objekte','avmztypen__objekte__SlgTeil').filter(schlagworte=id,)
   # typen = Muenztyp.objects.select_related('rv_bildtyp_set',).filter(rv_bildtyp__schlagworte=id,)
   # rvbildtypen = typen.values_list('rv_bildtyp__name', 'rv_bildtyp__id',).exclude(rv_bildtyp__name=None).distinct()
   # print(rvbildtypen.muenztyp.all())

   context = {
      'avbildtypen': avbildtypen,
   }
   
   return render(request, 'slg/avschlagwort_detail.html', context)

class AvSchlagwortList(generics.ListAPIView):
    serializer_class = AvSchlagwortSerializer

class AvBildtypSchlagwortList(generics.ListAPIView):
    serializer_class = AvBildtypSerializer

    def get_queryset(self):
        queryset = AvBildtyp.objects.all()
        bildtyp_query = self.request.query_params.get('bildtyp', None)
        schlagwort_query = self.request.query_params.get('schlagwort', None)

        if bildtyp_query:
            queryset = queryset.filter(
                Q(name__icontains=bildtyp_query) | Q(abk__icontains=bildtyp_query)
            )

        if schlagwort_query:
            queryset = queryset.filter(schlagworte__name__icontains=schlagwort_query)

        return queryset.distinct()

def RvSchlagwortView(request, id):
   
   # obj = Muenztyp.objects.select_related('Muenzstand', 'Herstellung', 'Nominal', 'Mzstaette', 'Metall', 'region', 'av_bildtyp', 'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen', 'Objekttyp',).prefetch_related('obj_set', 'mztyp_person_set', 'typ_ref_set',).get(id=id)
   # rvbildtypen = RvBildtyp.objects.prefetch_related('muenztyp_set','obj_set').exclude(name__isnull=True).filter(schlagworte=id,)
   rvbildtypen = RvBildtyp.objects.prefetch_related('rvmztypen','rvmztypen__objekte','rvmztypen__objekte__SlgTeil').filter(schlagworte=id,)
   # typen = Muenztyp.objects.select_related('rv_bildtyp_set',).filter(rv_bildtyp__schlagworte=id,)
   # rvbildtypen = typen.values_list('rv_bildtyp__name', 'rv_bildtyp__id',).exclude(rv_bildtyp__name=None).distinct()
   # print(rvbildtypen.muenztyp.all())
   context = {
      'rvbildtypen': rvbildtypen,
   }
   
   return render(request, 'slg/rvschlagwort_detail.html', context)

def RvSchlagwortTimelineView(request, id):
   
   # obj = Muenztyp.objects.select_related('Muenzstand', 'Herstellung', 'Nominal', 'Mzstaette', 'Metall', 'region', 'av_bildtyp', 'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen', 'Objekttyp',).prefetch_related('obj_set', 'mztyp_person_set', 'typ_ref_set',).get(id=id)
   # rvbildtypen = RvBildtyp.objects.prefetch_related('muenztyp_set','obj_set').exclude(name__isnull=True).filter(schlagworte=id,)
   rvbildtypen = RvBildtyp.objects.prefetch_related('rvmztypen','rvmztypen__objekte','rvmztypen__objekte__SlgTeil').filter(schlagworte=id,)
   startandenddate = rvbildtypen.aggregate(start=Min('rvmztypen__dat_von'), end=Max('rvmztypen__dat_bis'))
   
   # typen = Muenztyp.objects.select_related('rv_bildtyp_set',).filter(rv_bildtyp__schlagworte=id,)
   # rvbildtypen = typen.values_list('rv_bildtyp__name', 'rv_bildtyp__id',).exclude(rv_bildtyp__name=None).distinct()
   # print(rvbildtypen.muenztyp.all())
   context = {
      'rvbildtypen': rvbildtypen,
      'startandenddate': startandenddate,
   }
   
   return render(request, 'slg/timeline_bildtypen.html', context)

def PraegeherrTimelineView(request, id):
   
   rvbildtypen = RvBildtyp.objects.prefetch_related(
      'rvmztypen', 'rvmztypen__objekte', 'rvmztypen__objekte__SlgTeil'
   ).filter(rvmztypen__mztyp_person__idfk_Person=id).distinct()

   startandenddate = rvbildtypen.aggregate(start=Min('rvmztypen__dat_von'), end=Max('rvmztypen__dat_bis'))

   typen = Muenztyp.objects.filter(mztyp_person__idfk_Person=id).distinct()
   datierungen = typen.values('dat_von')

   context = {
      'rvbildtypen': rvbildtypen,
      'startandenddate': startandenddate,
      'mztyp': typen,
      'datierungen': datierungen,
   }
   
   return render(request, 'slg/timeline_praegeherr.html', context)

@api_view(['POST'])
@authentication_classes([QueryParamTokenAuthentication, SessionAuthentication, TokenAuthentication])
@permission_classes([IsAuthenticated])
def export_coins_api(request):
    """Accept a JSON payload {"invnrs": [...]} and return the coin list CSV export.

    Mirrors the behavior of the admin action "Münzliste exportieren" for remote callers
    (e.g. Concordia), but selects the Obj rows by inventory number instead of by admin
    queryset. Optionally accepts a `filename` field for the Content-Disposition.
    """
    payload = request.data or {}
    invnrs = payload.get('invnrs') or []
    if not isinstance(invnrs, list):
        return Response({'error': "'invnrs' must be a list"}, status=400)

    cleaned_invnrs = []
    seen = set()
    for value in invnrs:
        text = str(value or '').strip()
        if not text or text in seen:
            continue
        seen.add(text)
        cleaned_invnrs.append(text)

    if not cleaned_invnrs:
        return Response({'error': "'invnrs' list is empty"}, status=400)

    base_qs = Obj.objects.filter(invnr__in=cleaned_invnrs)
    distinct_pks = base_qs.values('pk').distinct().values_list('pk', flat=True)
    queryset = Obj.objects.filter(pk__in=distinct_pks).select_related(
        'Slg',
        'SlgTeil',
        'Typ',
        'idfk_Nominal',
        'Metall',
        'idfk_Mzstaette',
        'region',
        'av_bildtyp',
        'rv_bildtyp',
        'av_beizeichen',
        'rv_beizeichen',
        'av_offizin',
        'rv_offizin',
    ).prefetch_related(
        'Herstellungsmerkmale',
        'sekundaere_Merkmale',
        'Ppl',
        'idfk_Ref',
    )

    response = prepare_coin_data_export(queryset)

    filename = str(payload.get('filename') or '').strip() or 'exported_coins.csv'
    if '/' in filename or '\\' in filename or '\r' in filename or '\n' in filename:
        filename = 'exported_coins.csv'
    if not filename.lower().endswith('.csv'):
        filename = f"{filename}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # Informative header for proxies that need to know how many rows matched.
    response['X-Export-Matched-Invnrs'] = str(len(distinct_pks))
    response['X-Export-Requested-Invnrs'] = str(len(cleaned_invnrs))
    return response


def prepare_coin_data_export(queryset):
    # Stelle sicher, dass das queryset eindeutig ist
    queryset = queryset.distinct()
    
    import csv
    import io
    from django.http import StreamingHttpResponse
    from collections import defaultdict
    
    # Hilfsfunktion: str() nur wenn nicht None, sonst leerer String
    def safe_str(value):
        return str(value) if value is not None else ''
    
    # CSV Header definieren
    fieldnames = [
        'pk', 'workflow_objekt', 'workflow_mztyp', 'Inv.-Nr.', 'Durchmesser', 'Gewicht', 
        'Stempelstellung', 'Abnutzung', 'Sammlung', 'Sammlungsteil', 'Münzstätte', 'Region', 
        'praegeherren', 'Titel', 'Datierung verbale', 'Datierung von', 'Datierung bis', 
        'Objekttyp', 'Münzstand', 'Reichskreis', 'Nominal', 'Metall', 'herstellung', 
        'personen_av', 'avleg', 'av_bildtyp', 'av_schlagworte', 'av_beizeichen', 'av_offizin', 
        'personen_rv', 'rvleg', 'rv_bildtyp', 'rv_schlagworte', 'rv_beizeichen', 'rv_offizin', 
        'Wappen', 'Fälschung', 'Typ', 'Konkordanz', 'Zitate', 'Typ_unsicher', 'TempTyp', 'obj_ref_set', 
        'anmerkung', 'Herstellungsmerkmale', 'sekundaere_Merkmale', 'Avers', 'Revers',
        # Fundinformationen
        'fund_maßnahmennr', 'fund_fundnummer', 'fund_fundnummerzusatz', 'fund_kistennr',
        'fund_fundort', 'fund_parzelle', 'fund_fundstelle', 'fund_lm_von', 'fund_lm_bis',
        'fund_niveautiefe_von', 'fund_niveautiefe_bis', 'fund_niveau', 'fund_vn', 'fund_vno',
        'fund_vo', 'fund_vso', 'fund_vs', 'fund_vsw', 'fund_vw', 'fund_vnw', 'fund_schnitt',
        'fund_bereichsbezeichnung', 'fund_sondage', 'fund_quadrant', 'fund_quadrantzusatz',
        'fund_flaeche', 'fund_se', 'fund_stratum', 'fund_funddatum', 'fund_fundjahr',
        'fund_andere_materialien', 'fund_fundposition', 'fund_fundkontext', 'fund_anmerkung',
        'fund_bearbeiter'
    ]
    
    def data_generator():
        BATCH_SIZE = 500  # Prozessiere 500 Objekte auf einmal
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=fieldnames)
        
        # Header schreiben
        writer.writeheader()
        yield buffer.getvalue().encode('utf-8')
        buffer.seek(0)
        buffer.truncate(0)
        
        # Verbesserte Queryset-Optimierung
        optimized_queryset = queryset.select_related(
            'Slg', 'SlgTeil', 'Typ', 'Typ__Objekttyp', 'Typ__Herstellung', 
            'Typ__Reichskreis', 'Typ__Muenzstand', 'Typ__Nominal', 'Typ__Metall',
            'Typ__Mzstaette', 'Typ__region', 'Typ__av_bildtyp', 'Typ__rv_bildtyp',
            'Typ__av_beizeichen', 'Typ__rv_beizeichen', 'Typ__workflow',
            'idfk_Nominal', 'Metall', 'idfk_Mzstaette', 'region', 'workflow',
            'av_bildtyp', 'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen',
            'av_offizin', 'rv_offizin', 'faelschung', 'idfk_Muenzstand',
            'idfk_Herstellung', 'Objekttyp'
        ).prefetch_related(
            'Herstellungsmerkmale', 'sekundaere_Merkmale', 'obj_ref_set',
            'Typ__Konkordanz', 'Typ__mztyp_wappen_set__wappen'
        )
        
        # Prozessiere in Batches
        processed_count = 0
        seen_pks = set()
        
        for batch_start in range(0, queryset.count(), BATCH_SIZE):
            batch = list(optimized_queryset[batch_start:batch_start + BATCH_SIZE])
            
            if not batch:
                break
                
            # Alle IDs für diese Batch sammeln
            obj_ids = [obj.pk for obj in batch if obj.pk not in seen_pks]
            mztyp_ids = [obj.Typ.pk for obj in batch if obj.Typ and obj.pk not in seen_pks]
            
            # Bulk-Load aller zusätzlichen Daten
            # 1. Fund-Daten
            fund_data = {}
            if obj_ids:
                funds = Fund.objects.filter(objekt__pk__in=obj_ids).select_related('fundort')
                fund_data = {fund.objekt_id: fund for fund in funds}
            
            # 2. Schlagwort-Daten
            av_schlagworte_data = defaultdict(list)
            rv_schlagworte_data = defaultdict(list)
            if mztyp_ids:
                av_bildtyp_ids = [obj.Typ.av_bildtyp.pk for obj in batch if obj.Typ and obj.Typ.av_bildtyp]
                rv_bildtyp_ids = [obj.Typ.rv_bildtyp.pk for obj in batch if obj.Typ and obj.Typ.rv_bildtyp]
                
                if av_bildtyp_ids:
                    av_schlagworte = AvBildtyp_Schlagwort.objects.filter(
                        avbildtyp__pk__in=av_bildtyp_ids
                    ).select_related('schlagwort')
                    for rel in av_schlagworte:
                        av_schlagworte_data[rel.avbildtyp_id].append(str(rel.schlagwort))
                
                if rv_bildtyp_ids:
                    rv_schlagworte = RvBildtyp_Schlagwort.objects.filter(
                        rvbildtyp__pk__in=rv_bildtyp_ids
                    ).select_related('schlagwort')
                    for rel in rv_schlagworte:
                        rv_schlagworte_data[rel.rvbildtyp_id].append(str(rel.schlagwort))
            
            # 3. Personen-Daten (bulk)
            mztyp_personen_data = defaultdict(lambda: {'praegeherren': [], 'personen_av': [], 'personen_rv': []})
            obj_personen_data = defaultdict(lambda: {'praegeherren': [], 'personen_av': [], 'personen_rv': []})
            
            if mztyp_ids:
                mztyp_personen = Mztyp_Person.objects.filter(
                    Mztyp__pk__in=mztyp_ids
                ).select_related('idfk_Person', 'idfk_PersonFunktion')
                
                for rel in mztyp_personen:
                    person_name = rel.idfk_Person.name
                    if rel.idfk_PersonFunktion.id in [1, 6, 7]:
                        mztyp_personen_data[rel.Mztyp_id]['praegeherren'].append(person_name)
                    elif rel.idfk_PersonFunktion.id == 2:
                        if rel.appears_on_rev:
                            mztyp_personen_data[rel.Mztyp_id]['personen_rv'].append(person_name)
                        else:
                            mztyp_personen_data[rel.Mztyp_id]['personen_av'].append(person_name)
            
            if obj_ids:
                obj_personen = Obj_Person.objects.filter(
                    idfk_Obj__pk__in=obj_ids
                ).select_related('idfk_Person', 'idfk_PersonFunktion')
                
                for rel in obj_personen:
                    person_name = rel.idfk_Person.name
                    if rel.idfk_PersonFunktion.id in [1, 6, 7]:
                        obj_personen_data[rel.idfk_Obj_id]['praegeherren'].append(person_name)
                    elif rel.idfk_PersonFunktion.id == 2:
                        if rel.appears_on_rev:
                            obj_personen_data[rel.idfk_Obj_id]['personen_rv'].append(person_name)
                        else:
                            obj_personen_data[rel.idfk_Obj_id]['personen_av'].append(person_name)
            
            # Jetzt alle Objekte in diesem Batch verarbeiten
            rows_to_write = []
            
            for obj in batch:
                if obj.pk in seen_pks:
                    continue
                seen_pks.add(obj.pk)
                
                mztyp = obj.Typ
                fund = fund_data.get(obj.pk)
                
                # Personen-Daten
                if mztyp:
                    personen = mztyp_personen_data[mztyp.pk]
                else:
                    personen = obj_personen_data[obj.pk]
                
                # Schlagwort-Daten
                av_schlagworte = av_schlagworte_data.get(mztyp.av_bildtyp.pk if mztyp and mztyp.av_bildtyp else None, [])
                rv_schlagworte = rv_schlagworte_data.get(mztyp.rv_bildtyp.pk if mztyp and mztyp.rv_bildtyp else None, [])
                
                # Hilfswerte für Münzstätte und Region berechnen
                muenzstaette_val = obj.idfk_Mzstaette or (mztyp.Mzstaette if mztyp else None)
                region_val = obj.region or (mztyp.region if mztyp else None)
                av_bildtyp_val = obj.av_bildtyp or (mztyp.av_bildtyp if mztyp else None)
                rv_bildtyp_val = obj.rv_bildtyp or (mztyp.rv_bildtyp if mztyp else None)
                av_beizeichen_val = obj.av_beizeichen or (mztyp.av_beizeichen if mztyp else None)
                rv_beizeichen_val = obj.rv_beizeichen or (mztyp.rv_beizeichen if mztyp else None)
                
                data = {
                    'pk': obj.pk,
                    'workflow_objekt': safe_str(obj.workflow),
                    'workflow_mztyp': safe_str(mztyp.workflow) if mztyp else '',
                    'Inv.-Nr.': safe_str(obj.invnr),
                    'Durchmesser': obj.durchmesser,
                    'Gewicht': obj.gewicht,
                    'Stempelstellung': obj.stempelstellung,
                    'Abnutzung': obj.abnutzung,
                    'Sammlung': safe_str(obj.Slg),
                    'Sammlungsteil': safe_str(obj.SlgTeil),
                    'Münzstätte': safe_str(muenzstaette_val),
                    'Region': safe_str(region_val),
                    'praegeherren': ', '.join(personen['praegeherren']),
                    'Titel': obj.titel or (mztyp.titel if mztyp else '') or '',
                    'Datierung verbale': obj.dat_verb or (mztyp.dat_verb if mztyp else '') or '',
                    'Datierung von': obj.dat_von or (mztyp.dat_von if mztyp else None),
                    'Datierung bis': obj.dat_bis or (mztyp.dat_bis if mztyp else None),
                    'Objekttyp': safe_str(mztyp.Objekttyp if (mztyp and mztyp.Objekttyp) else obj.Objekttyp),
                    'Münzstand': safe_str(obj.idfk_Muenzstand or (mztyp.Muenzstand if mztyp else None)),
                    'Reichskreis': safe_str(mztyp.Reichskreis) if mztyp else '',
                    'Nominal': safe_str(obj.idfk_Nominal or (mztyp.Nominal if mztyp else None)),
                    'Metall': safe_str(obj.Metall or (mztyp.Metall if mztyp else None)),
                    'herstellung': safe_str(obj.idfk_Herstellung or (mztyp.Herstellung if mztyp else None) or 'Prägung'),
                    'personen_av': ', '.join(personen['personen_av']),
                    'avleg': obj.avleg or (mztyp.avleg if mztyp else '') or '',
                    'av_bildtyp': safe_str(av_bildtyp_val),
                    'av_schlagworte': ', '.join(av_schlagworte),
                    'av_beizeichen': safe_str(av_beizeichen_val),
                    'av_offizin': safe_str(obj.av_offizin),
                    'personen_rv': ', '.join(personen['personen_rv']),
                    'rvleg': obj.rvleg or (mztyp.rvleg if mztyp else '') or '',
                    'rv_bildtyp': safe_str(rv_bildtyp_val),
                    'rv_schlagworte': ', '.join(rv_schlagworte),
                    'rv_beizeichen': safe_str(rv_beizeichen_val),
                    'rv_offizin': safe_str(obj.rv_offizin),
                    'Wappen': ', '.join([str(w.wappen) for w in mztyp.mztyp_wappen_set.all()]) if mztyp else '',
                    'Fälschung': safe_str(obj.faelschung),
                    'Typ': safe_str(mztyp),
                    'Konkordanz': ', '.join([str(konkordanz) for konkordanz in mztyp.Konkordanz.all()]) if mztyp else '',
                    'Zitate': ('Referenz unsicher: ' if obj.Typ_unsicher else '') + (' = '.join([str(mztyp)] + [str(konkordanz) for konkordanz in mztyp.Konkordanz.all()]) if mztyp else (obj.anmerkung or '')),
                    'Typ_unsicher': obj.Typ_unsicher,
                    'TempTyp': obj.TempTyp or '',
                    'obj_ref_set': ', '.join([str(ref) for ref in obj.obj_ref_set.all()]),
                    'anmerkung': obj.anmerkung or '',
                    'Herstellungsmerkmale': ', '.join([str(merkmal) for merkmal in obj.Herstellungsmerkmale.all()]),
                    'sekundaere_Merkmale': ', '.join([str(merkmal) for merkmal in obj.sekundaere_Merkmale.all()]),
                }
                
                # Fundinformationen hinzufügen
                if fund:
                    data.update({
                        'fund_maßnahmennr': fund.maßnahmennr,
                        'fund_fundnummer': fund.fundnummer,
                        'fund_fundnummerzusatz': fund.fundnummerzusatz,
                        'fund_kistennr': fund.kistennr,
                        'fund_fundort': fund.fundort.name if fund.fundort else '',
                        'fund_parzelle': fund.parzelle,
                        'fund_fundstelle': fund.fundstelle,
                        'fund_lm_von': fund.lm_von,
                        'fund_lm_bis': fund.lm_bis,
                        'fund_niveautiefe_von': fund.niveautiefe_in_m_von,
                        'fund_niveautiefe_bis': fund.niveautiefe_in_m_bis,
                        'fund_niveau': fund.niveau,
                        'fund_vn': fund.vn, 'fund_vno': fund.vno, 'fund_vo': fund.vo,
                        'fund_vso': fund.vso, 'fund_vs': fund.vs, 'fund_vsw': fund.vsw,
                        'fund_vw': fund.vw, 'fund_vnw': fund.vnw,
                        'fund_schnitt': fund.schnitt,
                        'fund_bereichsbezeichnung': fund.bereichsbezeichnung,
                        'fund_sondage': fund.sondage, 'fund_quadrant': fund.quadrant,
                        'fund_quadrantzusatz': fund.quadrantzusatz, 'fund_flaeche': fund.flaeche,
                        'fund_se': fund.se, 'fund_stratum': fund.stratum,
                        'fund_funddatum': fund.funddatum, 'fund_fundjahr': fund.fundjahr,
                        'fund_andere_materialien': fund.andere_materialien,
                        'fund_fundposition': fund.fundposition, 'fund_fundkontext': fund.fundkontext,
                        'fund_anmerkung': fund.anmerkung, 'fund_bearbeiter': fund.bearbeiter,
                    })
                else:
                    # Setze alle Fund-Felder auf leeren String
                    for field in fieldnames:
                        if field.startswith('fund_') and field not in data:
                            data[field] = ''

                # Avers/Revers zusammensetzen
                av_components = []
                if data['avleg']:
                    av_components.append(data['avleg'])
                if data['av_bildtyp']:
                    av_components.append(data['av_bildtyp'])
                if data['av_beizeichen']:
                    beizeichen_text = data['av_beizeichen']
                    if "?" in beizeichen_text and data['av_offizin']:
                        beizeichen_text = beizeichen_text.replace("?", data['av_offizin'])
                    av_components.append(beizeichen_text)
                data['Avers'] = '. '.join(filter(None, av_components))
                
                rv_components = []
                if data['rvleg']:
                    rv_components.append(data['rvleg'])
                if data['rv_bildtyp']:
                    rv_components.append(data['rv_bildtyp'])
                if data['rv_beizeichen']:
                    beizeichen_text = data['rv_beizeichen']
                    if "?" in beizeichen_text and data['rv_offizin']:
                        beizeichen_text = beizeichen_text.replace("?", data['rv_offizin'])
                    rv_components.append(beizeichen_text)
                data['Revers'] = '. '.join(filter(None, rv_components))

                # String-Konvertierung für CSV-Kompatibilität
                for key, value in data.items():
                    if value is None:
                        data[key] = ''
                    elif not isinstance(value, (str, int, float)):
                        data[key] = str(value)

                rows_to_write.append(data)
            
            # Schreibe alle Zeilen dieses Batches auf einmal
            for row_data in rows_to_write:
                writer.writerow(row_data)
            
            # Yield den gesamten Batch
            if rows_to_write:
                yield buffer.getvalue().encode('utf-8')
                buffer.seek(0)
                buffer.truncate(0)

    response = StreamingHttpResponse(
        data_generator(),
        content_type='text/csv; charset=utf-8'
    )
    response['Content-Disposition'] = 'attachment; filename="exported_coins.csv"'
    return response

def prepare_coin_data(obj):
    mztyp = obj.Typ
    
    # Use prefetched data instead of new queries
    av_bildtyp_obj = obj.av_bildtyp or (mztyp.av_bildtyp if mztyp else None)
    rv_bildtyp_obj = obj.rv_bildtyp or (mztyp.rv_bildtyp if mztyp else None)

    # Schlagworte from prefetched sets
    av_schlagworte_qs = av_bildtyp_obj.avbildtyp_schlagwort_set.all() \
        if av_bildtyp_obj and hasattr(av_bildtyp_obj, 'avbildtyp_schlagwort_set') else []
    
    rv_schlagworte_qs = rv_bildtyp_obj.rvbildtyp_schlagwort_set.all() \
        if rv_bildtyp_obj and hasattr(rv_bildtyp_obj, 'rvbildtyp_schlagwort_set') else []

    # Personen from prefetched sets
    mztyp_personen = mztyp.mztyp_person_set.all() \
        if mztyp and hasattr(mztyp, 'mztyp_person_set') else []
    obj_personen = obj.obj_person_set.all() if hasattr(obj, 'obj_person_set') else []
    
    # Wappen from prefetched set
    wappen_qs = mztyp.mztyp_wappen_set.all() \
        if mztyp and hasattr(mztyp, 'mztyp_wappen_set') else None

    data = {
         'pk': obj.pk,
         'invnr': obj.invnr,
         'durchmesser': obj.durchmesser,
         'gewicht': obj.gewicht,
         'stempelstellung': obj.stempelstellung,
         'av_beizeichen': obj.av_beizeichen or (mztyp.av_beizeichen if mztyp else None),
         'rv_beizeichen': obj.rv_beizeichen or (mztyp.rv_beizeichen if mztyp else None),
         'av_offizin': obj.av_offizin,
         'rv_offizin': obj.rv_offizin,
         'av_offizin_symbol': obj.av_offizin_symbol or (mztyp.av_offizin_symbol if mztyp else None),
         'rv_offizin_symbol': obj.rv_offizin_symbol or (mztyp.rv_offizin_symbol if mztyp else None),
         'bild_urls': obj.get_bild_urls(),
         'titel': obj.titel or (mztyp.titel if mztyp else None),
         'dat_verb': obj.dat_verb or (mztyp.dat_verb if mztyp else None),
         'objekttyp': mztyp.Objekttyp if (mztyp and mztyp.Objekttyp) else (obj.Objekttyp or None),
         'muenzstand': obj.idfk_Muenzstand or (mztyp.Muenzstand if mztyp else None),
         'reichskreis': mztyp.Reichskreis if mztyp else None,
         'nominal': obj.idfk_Nominal or (mztyp.Nominal if mztyp else None),
         'metall': obj.Metall or (mztyp.Metall if mztyp else None),
         'herstellung': obj.idfk_Herstellung or (mztyp.Herstellung if mztyp else None),
         'avleg': obj.avleg or (mztyp.avleg if mztyp else None),
         'rvleg': obj.rvleg or (mztyp.rvleg if mztyp else None),
         'av_bildtyp': obj.av_bildtyp or (mztyp.av_bildtyp if mztyp else None),
         'rv_bildtyp': obj.rv_bildtyp or (mztyp.rv_bildtyp if mztyp else None),
         'av_schlagworte': av_schlagworte_qs,
         'rv_schlagworte': rv_schlagworte_qs,
         'personen': list(chain(mztyp_personen, obj_personen)),
         'wappen': wappen_qs,
         'slg': obj.Slg,
         'slgteil': obj.SlgTeil,
         'mzstaette': obj.idfk_Mzstaette or (mztyp.Mzstaette if mztyp else None),
         'av_bildrand': obj.av_bildrand or (mztyp.av_bildrand if mztyp else None),
         'rv_bildrand': obj.rv_bildrand or (mztyp.rv_bildrand if mztyp else None),
         'faelschung': obj.faelschung,
         'region': obj.region or (mztyp.region if mztyp else None),
         'Typ': mztyp,
         'Konkordanz': mztyp.Konkordanz.all() if mztyp else None,
         'idfk_Ref': obj.idfk_Ref,
         'Typ_unsicher': obj.Typ_unsicher,
         'obj_ref_set': list(obj.obj_ref_set.all()),
         'anmerkung': obj.anmerkung,
         'Herstellungsmerkmale': list(obj.Herstellungsmerkmale.all()),
         'sekundaere_Merkmale': list(obj.sekundaere_Merkmale.all()),
         'av_beizeichen_display': None,
         'rv_beizeichen_display': None,
         'av_offizin_display': None,
         'rv_offizin_display': None,
         'av_offizin_symbol_html': None,
         'rv_offizin_symbol_html': None,
         'auflage': mztyp.auflage if mztyp else None,
         'beschlussdatum': mztyp.beschlussdatum if mztyp else None,
         'ausgabedatum': mztyp.ausgabedatum if mztyp else None,
    }
    
    # Prepare offizin symbol HTML if available
    if data['av_offizin_symbol']:
        data['av_offizin_symbol_html'] = data['av_offizin_symbol'].get_svg_html()
    
    if data['rv_offizin_symbol']:
        data['rv_offizin_symbol_html'] = data['rv_offizin_symbol'].get_svg_html()

    # Anpassung der Beizeichen-Anzeige für Vorderseite
    if data['av_beizeichen']:
        beizeichen_text = data['av_beizeichen'].name
        
        # Replace [Offizinszeichen] and [Monogramm] with SVG if available
        if data['av_offizin_symbol_html']:
            if "[Offizinszeichen]" in beizeichen_text:
                beizeichen_text = beizeichen_text.replace("[Offizinszeichen]", data['av_offizin_symbol_html'])
            if "[Monogramm]" in beizeichen_text:
                beizeichen_text = beizeichen_text.replace("[Monogramm]", data['av_offizin_symbol_html'])
        
        # Replace specific monogram references like [Lorber Monogram 219]
        beizeichen_text = replace_monogram_references(beizeichen_text)
        
        # Handle the existing ? replacement logic
        if "?" in beizeichen_text and data['av_offizin']:
            beizeichen_text = beizeichen_text.replace("?", data['av_offizin'].name)
            data['av_offizin_display'] = None  # Setze auf None, da es bereits im Beizeichen enthalten ist
        else:
            data['av_offizin_display'] = data['av_offizin'].name if data['av_offizin'] else None
            
        data['av_beizeichen_display'] = beizeichen_text
    elif data['av_offizin']:
        data['av_beizeichen_display'] = None
        data['av_offizin_display'] = data['av_offizin'].name

    # Ähnliche Logik für die Rückseite
    if data['rv_beizeichen']:
        beizeichen_text = data['rv_beizeichen'].name
        
        # Replace [Offizinszeichen] and [Monogramm] with SVG if available
        if data['rv_offizin_symbol_html']:
            if "[Offizinszeichen]" in beizeichen_text:
                beizeichen_text = beizeichen_text.replace("[Offizinszeichen]", data['rv_offizin_symbol_html'])
            if "[Monogramm]" in beizeichen_text:
                beizeichen_text = beizeichen_text.replace("[Monogramm]", data['rv_offizin_symbol_html'])
        
        # Replace specific monogram references like [Lorber Monogram 219]
        beizeichen_text = replace_monogram_references(beizeichen_text)
        
        # Handle the existing ? replacement logic
        if "?" in beizeichen_text and data['rv_offizin']:
            beizeichen_text = beizeichen_text.replace("?", data['rv_offizin'].name)
            data['rv_offizin_display'] = None
        else:
            data['rv_offizin_display'] = data['rv_offizin'].name if data['rv_offizin'] else None
            
        data['rv_beizeichen_display'] = beizeichen_text
    elif data['rv_offizin']:
        data['rv_beizeichen_display'] = None
        data['rv_offizin_display'] = data['rv_offizin'].name

    return data

def ObjektView(request, id):
    
    # Custom prefetch for schlagworte to avoid N+1 queries
    av_schlagwort_prefetch = Prefetch(
        'av_bildtyp__avbildtyp_schlagwort_set',
        queryset=AvBildtyp_Schlagwort.objects.select_related('schlagwort')
    )
    
    rv_schlagwort_prefetch = Prefetch(
        'rv_bildtyp__rvbildtyp_schlagwort_set', 
        queryset=RvBildtyp_Schlagwort.objects.select_related('schlagwort')
    )
    
    # Custom prefetch for Typ's av/rv bildtyp schlagworte
    typ_av_schlagwort_prefetch = Prefetch(
        'Typ__av_bildtyp__avbildtyp_schlagwort_set',
        queryset=AvBildtyp_Schlagwort.objects.select_related('schlagwort')
    )
    
    typ_rv_schlagwort_prefetch = Prefetch(
        'Typ__rv_bildtyp__rvbildtyp_schlagwort_set',
        queryset=RvBildtyp_Schlagwort.objects.select_related('schlagwort') 
    )
    
    # Custom prefetch for person relationships
    obj_person_prefetch = Prefetch(
        'obj_person_set',
        queryset=Obj_Person.objects.select_related('idfk_Person', 'idfk_PersonFunktion').order_by('idfk_PersonFunktion__name')
    )
    
    mztyp_person_prefetch = Prefetch(
        'Typ__mztyp_person_set',
        queryset=Mztyp_Person.objects.select_related('idfk_Person', 'idfk_PersonFunktion').order_by('idfk_PersonFunktion__name')
    )
    
    # Custom prefetch for wappen
    mztyp_wappen_prefetch = Prefetch(
        'Typ__mztyp_wappen_set',
        queryset=Mztyp_Wappen.objects.select_related('wappen').order_by('wappen__name')
    )

    obj_ref_prefetch = Prefetch(
        'obj_ref_set',
        queryset=Obj_Ref.objects.select_related('idfk_Ref', 'variante')
    )

    obj_queryset = Obj.objects.select_related(
        # Existing relationships
        'idfk_Muenzstand', 'idfk_Herstellung', 'Slg', 'SlgTeil', 'idfk_Nominal', 
        'idfk_Mzstaette', 'Metall', 'Typ', 'faelschung', 'region', 'av_bildtyp', 
        'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen', 'av_bildrand', 'rv_bildrand', 
        'Objekttyp',
        
        # Missing relationships that cause extra queries
        'workflow',
        'av_offizin', 'rv_offizin',
        'av_offizin_symbol', 'rv_offizin_symbol',
        'rand',
        
        # Typ relationships to reduce queries
        'Typ__workflow',
        'Typ__Objekttyp', 'Typ__Herstellung', 'Typ__Muenzstand', 'Typ__Nominal',
        'Typ__Metall', 'Typ__Mzstaette', 'Typ__region', 'Typ__av_bildtyp',
        'Typ__rv_bildtyp', 'Typ__av_beizeichen', 'Typ__rv_beizeichen',
        'Typ__av_bildrand', 'Typ__rv_bildrand', 'Typ__rand',
        'Typ__av_offizin_symbol', 'Typ__rv_offizin_symbol',
        'Typ__auflage', 'Typ__Ref'
        
    ).prefetch_related(
        'Herstellungsmerkmale', 'sekundaere_Merkmale',
        obj_ref_prefetch,
        mztyp_person_prefetch,
        
        av_schlagwort_prefetch,
        rv_schlagwort_prefetch,
        typ_av_schlagwort_prefetch,
        typ_rv_schlagwort_prefetch,
        
        mztyp_wappen_prefetch,
        
        'Typ__schlagworte',
        'idfk_Ref'
        
    )

    obj = get_object_or_404(obj_queryset, id=id)
    
    coin_data = prepare_coin_data(obj)
    
    context = {
        'coin_data': coin_data,
        'test': obj,
    }
    
    return render(request, 'slg/details.html', context)


def objekt_detail_partial(request, id):
    obj = get_object_or_404(Objekt, id=id)
    return render(request, 'objekt_detail_partial.html', {'Obj': obj})


def kontakt(request):
    if request.method == 'POST':
        name = request.POST.get('name', '')
        email = request.POST.get('email', '')
        betreff = request.POST.get('betreff', '')
        nachricht = request.POST.get('nachricht', '')
        
        # Kontaktanfrage in der Datenbank speichern
        try:
            Kontaktanfrage.objects.create(
                name=name,
                email=email,
                betreff=betreff,
                nachricht=nachricht
            )
            messages.success(request, 'Ihre Nachricht wurde erfolgreich gesendet. Wir werden uns so bald wie möglich bei Ihnen melden.')
        except Exception as e:
            messages.error(request, 'Beim Speichern Ihrer Nachricht ist ein Fehler aufgetreten. Bitte versuchen Sie es später erneut.')
            print(f"Kontaktformular-Fehler: {e}")
        
        # Zurück zur Startseite mit Anker zum Kontaktformular
        return redirect('/#kontakt')
    
    # Bei GET-Anfragen einfach zur Startseite weiterleiten
    return redirect('/')

def objekt_aenderung(request, objekt_id):
    objekt = get_object_or_404(Obj, id=objekt_id)
    
    if request.method == 'POST':
        name = request.POST.get('name', '')
        email = request.POST.get('email', '')
        feld = request.POST.get('feld', '')
        alter_wert = request.POST.get('alter_wert', '')
        neuer_wert = request.POST.get('neuer_wert', '')
        begruendung = request.POST.get('begruendung', '')
        
        # Änderungsvorschlag speichern
        try:
            # Prüfen, ob alle erforderlichen Felder vorhanden sind
            if not all([name, email, feld, neuer_wert, begruendung]):
                raise ValueError("Nicht alle erforderlichen Felder wurden ausgefüllt.")
            
            # Änderungsvorschlag erstellen
            aenderung = ObjektAenderung.objects.create(
                objekt=objekt,
                name=name,
                email=email,
                feld=feld,
                alter_wert=alter_wert,
                neuer_wert=neuer_wert,
                begruendung=begruendung
            )
            
            # Erfolgreiche Erstellung des Änderungsvorschlags
            messages.success(request, 'Ihr Änderungsvorschlag wurde erfolgreich übermittelt und wird geprüft.')
            
            # Optional: Benachrichtigung an Administratoren senden
            # send_admin_notification(aenderung)
            
        except Exception as e:
            # Detaillierte Fehlerinformationen für die Fehlersuche
            error_details = traceback.format_exc()
            print(f"Fehler beim Speichern des Änderungsvorschlags: {e}")
            print(f"Fehlerdetails: {error_details}")
            
            # Benutzerfreundliche Fehlermeldung
            messages.error(request, f'Beim Speichern Ihres Änderungsvorschlags ist ein Fehler aufgetreten: {str(e)}')
        
        # Zurück zur Detailseite
        return HttpResponseRedirect(f'/objekt/{objekt_id}/')
    
    # Bei GET-Anfragen zur Detailseite zurückleiten
    return HttpResponseRedirect(f'/objekt/{objekt_id}/')

# Fügen Sie diese Funktion zu Ihrer views.py hinzu
def about(request):
    return render(request, 'slg/about.html')

def impressum(request):
    return render(request, 'slg/impressum.html')

def datenschutz(request):
    return render(request, 'slg/datenschutz.html')

def graph_view(request):
    """View function to display the graph visualization tool."""
    return render(request, 'slg/graphv0.2.html')

def _get_filtered_mtoa_queryset(request):
    """
    Zentrale Browse-Filterlogik für MTOA, damit Liste, Chart und andere
    Auswertungen dieselbe Treffermenge verwenden.
    """
    MTOA_FILTERS = {
        'q': {'fields': ['invnr', 'objekttitel', 'rv_legende', 'av_legende', 'typ'], 'lookup': 'icontains'},
        'avleg': {'fields': ['av_legende'], 'lookup': 'icontains'},
        'rvleg': {'fields': ['rv_legende'], 'lookup': 'icontains'},
        'Muenzstaette': {'fields': ['mzstaette'], 'lookup': 'exact'},
        'Muenzstand': {'fields': ['muenzstand'], 'lookup': 'exact'},
        'Reichskreis': {'fields': ['reichskreis'], 'lookup': 'exact'},
        'region': {'fields': ['region'], 'lookup': 'exact'},
        'Nominal': {'fields': ['nominal'], 'lookup': 'exact'},
        'Nominal_id': {'fields': ['nominal_fk_id'], 'lookup': 'in'},
        'material': {'fields': ['metall'], 'lookup': 'icontains'},
        'Slg': {'fields': ['slg_fk_id'], 'lookup': 'exact'},
        'SlgTeil': {'fields': ['slgteil_fk_id'], 'lookup': 'exact'},
        'av_bildtyp': {'fields': ['av_bildtyp_fk_id'], 'lookup': 'exact'},
        'av_beizeichen': {'fields': ['av_beizeichen'], 'lookup': 'exact'},
        'rv_bildtyp': {'fields': ['rv_bildtyp_fk_id'], 'lookup': 'exact'},
        'rv_beizeichen': {'fields': ['rv_beizeichen'], 'lookup': 'exact'},
        'rv_schlagwort': {'fields': ['rv_schlagworte'], 'lookup': 'icontains'},
        'av_schlagwort': {'fields': ['av_schlagworte'], 'lookup': 'icontains'},
        'obj_type': {'fields': ['objekttyp'], 'lookup': 'exact'},
        'objekttyp': {'fields': ['objekttyp_fk_id'], 'lookup': 'exact'},
        'coin_type': {'fields': ['typ_fk_id'], 'lookup': 'in'},
    }

    unbestimmt_param = request.GET.get('unbestimmt', '')
    if unbestimmt_param == 'True':
        is_unbestimmt = True
        qs = MuenztypObjektAnzeige.objects.filter(typ_fk__isnull=True)
    elif unbestimmt_param == 'False':
        is_unbestimmt = False
        qs = MuenztypObjektAnzeige.objects.filter(typ_fk__isnull=False)
    else:
        is_unbestimmt = False
        qs = MuenztypObjektAnzeige.objects.all()

    general_filters = []
    needs_distinct = False

    for param, config in MTOA_FILTERS.items():
        values = [v for v in request.GET.getlist(param) if is_valid_qparam(v)]
        if not values:
            continue

        if config['lookup'] == 'in':
            general_filters.append(Q(**{f"{config['fields'][0]}__in": values}))
        else:
            per_value = [
                Q(**{f"{field}__{config['lookup']}": val})
                for val in values
                for field in config['fields']
            ]
            general_filters.append(reduce(or_, per_value))

    date_from = request.GET.get('dat_von')
    date_to = request.GET.get('dat_bis')
    if is_valid_qparam(date_from) and is_valid_qparam(date_to):
        general_filters.append(
            Q(datierung_von__lte=date_to) &
            Q(datierung_bis__gte=date_from)
        )

    if general_filters:
        qs = qs.filter(reduce(and_, general_filters))

    person_params = {
        'Praegeherren': {'funktion_ids': [1, 6, 7], 'appears_on_rev': None},
        'Dargestellte_AV': {'funktion_ids': [2], 'appears_on_rev': False},
        'Dargestellte_RV': {'funktion_ids': [2], 'appears_on_rev': True},
        'Person': {'exclude_funktion_ids': [1, 2, 6, 7], 'appears_on_rev': None},
    }

    for param_name, config in person_params.items():
        if param_name in request.GET:
            names = [name for name in request.GET.getlist(param_name) if is_valid_qparam(name)]
            if names:
                needs_distinct = True
                name_q = Q(mtoaperson__person__name__in=names)

                if config.get('funktion_ids'):
                    name_q &= Q(mtoaperson__funktion_id__in=config['funktion_ids'])
                elif config.get('exclude_funktion_ids'):
                    name_q &= ~Q(mtoaperson__funktion_id__in=config['exclude_funktion_ids'])

                if config.get('appears_on_rev') is not None:
                    name_q &= Q(mtoaperson__appears_on_rev=config['appears_on_rev'])

                qs = qs.filter(name_q)

    wappen_names = [name for name in request.GET.getlist('Wappen') if is_valid_qparam(name)]
    if wappen_names:
        needs_distinct = True
        qs = qs.filter(typ_fk__wappen__name__in=wappen_names)

    ref_values = [name for name in request.GET.getlist('Ref') if is_valid_qparam(name)]
    if ref_values:
        ref_query = Q()
        if unbestimmt_param != 'True':
            ref_query |= Q(Typ__Ref__abk__in=ref_values)
        if unbestimmt_param != 'False':
            ref_query |= Q(idfk_Ref__abk__in=ref_values)

        ref_obj_ids = (
            Obj.objects.filter(ref_query)
            .values_list('id', flat=True)
            .distinct()
        )
        qs = qs.filter(obj_id__in=ref_obj_ids)

    her_merk_values = [name for name in request.GET.getlist('her_merk') if is_valid_qparam(name)]
    if her_merk_values:
        obj_ids = set(
            Obj.objects.filter(Herstellungsmerkmale__name__in=her_merk_values)
            .values_list('id', flat=True)
        )
        qs = qs.filter(obj_id__in=obj_ids)

    sek_merk_values = [name for name in request.GET.getlist('sek_merk') if is_valid_qparam(name)]
    if sek_merk_values:
        obj_ids = set(
            Obj.objects.filter(sekundaere_Merkmale__name__in=sek_merk_values)
            .values_list('id', flat=True)
        )
        qs = qs.filter(obj_id__in=obj_ids)

    if needs_distinct:
        qs = qs.distinct()

    qs = qs.annotate(
        _unbestimmt_sort=Case(
            When(typ_fk__isnull=True, then=Value(1)),
            default=Value(0),
            output_field=IntegerField(),
        )
    )
    return qs.order_by('_unbestimmt_sort', 'datierung_von', 'datierung_bis', 'pk'), is_unbestimmt, needs_distinct

# Neuer Endpoint für Chart-Daten
def area_chart_data(request):
    try:
        print(f"DEBUG: Request GET params: {request.GET}")

        filtered_qs, _, _ = _get_filtered_mtoa_queryset(request)
        obj_ids = list(filtered_qs.values_list('obj_id', flat=True))
        print(f"DEBUG: Gefundene MTOA-Objekte: {len(obj_ids)}")

        if not obj_ids:
            return JsonResponse({
                'years': [],
                'counts': [],
                'percentages': [],
                'debug': {'message': 'Keine Objekte mit passenden Filtern gefunden'}
            })

        pairs = list(
            filtered_qs.exclude(
                Q(datierung_von=None) | Q(datierung_bis=None)
            ).values_list('datierung_von', 'datierung_bis')
        )
        print(f"DEBUG: MTOA-Datumswerte: {len(pairs)} Paare gefunden")

        # Wenn Datierungsfilter aktiv sind, die Zeitspannen für den Chart
        # auf den angeforderten Bereich begrenzen. Sonst bleibt der Ausschnitt
        # durch lange Objektspannen unnötig weit.
        requested_from = request.GET.get('dat_von')
        requested_to = request.GET.get('dat_bis')
        chart_from = None
        chart_to = None

        try:
            if is_valid_qparam(requested_from):
                chart_from = int(requested_from)
            if is_valid_qparam(requested_to):
                chart_to = int(requested_to)
        except (TypeError, ValueError):
            chart_from = None
            chart_to = None

        if chart_from is not None and chart_to is not None and chart_from > chart_to:
            chart_from, chart_to = chart_to, chart_from

        print(f"DEBUG: chart_from={chart_from}, chart_to={chart_to}")

        max_span = 3000
        if chart_from is not None and chart_to is not None:
            max_span = chart_to - chart_from + 2

        years_counter = Counter()
        error_count = 0

        for start, end in pairs:
            try:
                start = int(start)
                end = int(end)

                if chart_from is not None:
                    start = max(start, chart_from)
                if chart_to is not None:
                    end = min(end, chart_to)

                if end < start:
                    continue

                span = end - start + 1

                if span <= 0 or span > max_span:
                    error_count += 1
                    continue

                weight = 1.0 / span
                for y in range(start, end + 1):
                    years_counter[y] += weight
            except (TypeError, ValueError) as e:
                error_count += 1
                continue

        print(f"DEBUG: Fehlerhafte Datumsbereiche: {error_count}")
        print(f"DEBUG: Verarbeitete Jahre im Counter: {len(years_counter)}")
        
        # Vollständige Jahresliste
        if years_counter:
            try:
                min_year = min(years_counter.keys())
                max_year = max(years_counter.keys())
                years = list(range(min_year, max_year + 1))
                counts = [years_counter.get(y, 0) for y in years]
            except ValueError:
                years = []
                counts = []
        else:
            years = []
            counts = []
        
        # Prozent-Rechnung
        total = sum(counts)
        if total > 0:
            percentages = [round(c / total * 100, 1) for c in counts]
        else:
            percentages = [0 for _ in counts]
        
        return JsonResponse({
            'years': years,
            'counts': counts,
            'percentages': percentages,
            'debug': {
                'filtered_objects': len(obj_ids),
                'date_pairs': len(pairs),
                'error_count': error_count,
                'years_counter_items': len(years_counter),
                'chart_from': chart_from,
                'chart_to': chart_to,
                'years_min': years[0] if years else None,
                'years_max': years[-1] if years else None,
            }
        })
    except Exception as e:
        import traceback
        error_details = traceback.format_exc()
        logger.error(f"ERROR in area_chart_data: {error_details}")

        return JsonResponse({
            'error': 'Ungültige oder inkompatible Filterparameter.',
            'years': [],
            'counts': [],
            'percentages': []
        }, status=400)






# In views.py

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen
def facet_api(request):
    facet_field = request.GET.get('facet')
    if not facet_field:
        return Response({"error": "Parameter 'facet' ist erforderlich"}, status=400)

    unbestimmt_param = request.GET.get('unbestimmt', '')
    context_key = 'unbestimmt' if unbestimmt_param == 'True' else 'default'
    # Lade die gesamte Konfiguration für den Kontext
    all_configs = FILTER_PARAMETERS.get(context_key, {})
    config = all_configs.get(facet_field)

    if not config:
        # Prüfen, ob es ein Personen-Facet mit spezifischer Konfiguration ist
        person_configs = {k: v for k, v in all_configs.items() if k in ['Praegeherren', 'Dargestellte_AV', 'Dargestellte_RV', 'Person']}
        if facet_field in person_configs:
            config = person_configs[facet_field]
        else:
             return Response({"error": f"Unbekanntes Facetten-Feld oder fehlende Konfiguration: {facet_field}"}, status=400)


    # ------------- Feld-Definition (weitgehend unverändert) -----------------
    field_path = config.get('facet_field', config['fields'][0])
    id_field = config.get('id_field')
    if not id_field and '__name' in field_path:
        id_field = field_path.replace('__name', '__id')
    # Speziell für Personenpfade den Basispfad bestimmen
    is_person_facet = facet_field in ['Praegeherren', 'Dargestellte_AV', 'Dargestellte_RV', 'Person']
    person_base = None
    if is_person_facet:
         person_base = 'obj_person' if context_key == 'unbestimmt' else 'Typ__mztyp_person'
         # Sicherstellen, dass field_path und id_field korrekt sind für den Kontext
         field_path = f"{person_base}__idfk_Person__name"
         id_field = f"{person_base}__idfk_Person_id"


    # ----------- Basismenge (Obj oder Typ bestimmt / unbestimmt) ----------
    if unbestimmt_param == 'True':
        base_qs = Obj.objects.filter(Typ__isnull=True)
    elif unbestimmt_param == 'False':
        base_qs = Obj.objects.filter(Typ__isnull=False)
    else:  # 'Alle' or empty
        base_qs = Obj.objects.all()
    # -----------------------------------------------------------------------

    # ----------- alle anderen URL-Filter anwenden (mit Facet-Kontext) -----
    # Hier werden andere Personenfilter nur nach Namen angewendet
    qs = apply_filters(request, base_qs, exclude_field=facet_field, facet_context=True)
    # -----------------------------------------------------------------------

    # ----------- *Zusätzlicher* spezifischer Filter für das *aktuelle* Personen-Facet --
    if is_person_facet and person_base:
        facet_filter = Q() # Leeres Q-Objekt
        function_ids = config.get('person_function_ids')
        exclude_ids = config.get('exclude_function_ids')
        appears_on_rev = config.get('appears_on_rev')

        if function_ids:
             facet_filter &= Q(**{f'{person_base}__idfk_PersonFunktion_id__in': function_ids})
        elif exclude_ids:
             facet_filter &= ~Q(**{f'{person_base}__idfk_PersonFunktion_id__in': exclude_ids})

        if appears_on_rev is not None:
            facet_filter &= Q(**{f'{person_base}__appears_on_rev': appears_on_rev})

        if facet_filter: # Nur filtern, wenn Bedingungen vorhanden sind
             qs = qs.filter(facet_filter)
    # ---------------------------------------------------------------------------

    # ----------- freie Suchphrase (Select2 "term" oder "q") ---------------
    search_term = request.GET.get('term') or request.GET.get('q')
    if search_term:
        # Wichtig: Suche im korrekten Namensfeld (besonders für Personen)
        search_field = field_path # field_path sollte jetzt korrekt sein
        qs = qs.filter(**{f"{search_field}__icontains": search_term})
    # -----------------------------------------------------------------------

    # ----------- leere Felder raus, zählen, sortieren ---------------------
    values_dict = {'value': F(field_path)}
    if id_field:
         values_dict['obj_id'] = F(id_field)

    # --- Exclude NULL and potentially empty strings ---
    exclude_empty_filter = {f"{field_path}__isnull": True}

    # Heuristik: Wenn der Pfad auf '__name' oder ähnliche Textfelder endet,
    # schließe zusätzlich leere Strings aus. Passe die Endungen bei Bedarf an.
    if field_path.endswith('__name') or field_path.endswith('__bezeichnung') or field_path.endswith('__label'): # Füge ggf. weitere hinzu
        # Prüfe zur Sicherheit, ob der Key nicht schon existiert (sollte nicht, aber schadet nicht)
        if f"{field_path}__exact" not in exclude_empty_filter:
            exclude_empty_filter[f"{field_path}__exact"] = ''

    qs = (
        qs.exclude(**exclude_empty_filter)
          .values(**values_dict)
          .annotate(count=Count('pk', distinct=True))
          .order_by('value' if search_term else '-count')
    )

    # Limit erst nach der Aggregation anwenden
    qs = qs[:50]
    # -----------------------------------------------------------------------

    # ----------- JSON-Antwort (weitgehend unverändert) ----------------------
    result = []
    for entry in qs:
        # Erneut prüfen, ob 'value' leer ist, sicherheitshalber
        if entry['value']:
            obj = {'name': entry['value'], 'count': entry['count']}
            if 'obj_id' in entry:
                obj['id'] = entry['obj_id']
            result.append(obj)

    return Response(result)


@api_view(['GET'])
@authentication_classes([QueryParamTokenAuthentication, SessionAuthentication, TokenAuthentication])
@permission_classes([IsAuthenticated])
def get_konkordanzen(request):
    typ_id = request.GET.get('typ', None)
    konkordanzen = []

    if not typ_id:
        return JsonResponse({"konkordanzen": konkordanzen})

    try:
        typ_id = int(typ_id)
    except (TypeError, ValueError):
        return JsonResponse({"error": "Parameter 'typ' muss eine Zahl sein.", "konkordanzen": []}, status=400)

    typ = Muenztyp.objects.filter(id=typ_id).first()
    if not typ:
        return JsonResponse({"error": "Muenztyp nicht gefunden.", "konkordanzen": []}, status=404)

    for konkordanz in typ.Konkordanz.all():
        konkordanzen.append({"id": konkordanz.id, "titel": str(konkordanz)})

    return JsonResponse({"konkordanzen": konkordanzen})

@api_view(['POST'])
@authentication_classes([QueryParamTokenAuthentication, SessionAuthentication, TokenAuthentication])
@permission_classes([IsAuthenticated])
def konkordanzen_bulk(request):
    raw_type_ids = request.data.get('type_ids') or request.data.get('ids') or []
    if not isinstance(raw_type_ids, list):
        return Response({'ok': False, 'error': 'type_ids must be a list'}, status=400)

    type_ids = []
    seen = set()
    for raw in raw_type_ids:
        try:
            type_id = int(raw)
        except (TypeError, ValueError):
            continue
        if type_id <= 0 or type_id in seen:
            continue
        seen.add(type_id)
        type_ids.append(type_id)

    if not type_ids:
        return Response({'ok': True, 'konkordanzen_by_type': {}, 'errors': {}})

    konkordanzen_by_type = {str(type_id): [] for type_id in type_ids}
    qs = Muenztyp.objects.filter(id__in=type_ids).prefetch_related('Konkordanz')
    for typ in qs:
        konkordanzen_by_type[str(typ.id)] = [
            {'id': konkordanz.id, 'titel': str(konkordanz)}
            for konkordanz in typ.Konkordanz.all()
        ]

    return Response({'ok': True, 'konkordanzen_by_type': konkordanzen_by_type, 'errors': {}})

class MzstaettenRView(viewsets.ModelViewSet):
# class MzstaettenRView(generics.RetrieveUpdateDestroyAPIView):
    
    queryset         = Mzstaette.objects.all()
    serializer_class = MzstaettenSerializer
    

   #  def get_queryset(self):
   #      return Mzstaette.objects.all()
    
    # def get_object(self):
    #     id = self.kwargs.get("id")
    #     return Mzstaette.objects.get(id=id)

def jsonresp(request, id):
   # obj = Obj.objects.select_related('idfk_Muenzstand', 'idfk_Herstellung', 'idfk_SlgTeil', 'idfk_Nominal', 'idfk_Mzstaette').prefetch_related('Ppl').get(id=id)
   obj = Obj.objects.get(pk=id)
   serializer = ObjSerializer(obj)
   return JsonResponse(serializer.data)
   #data = serializers.serialize(obj, many=True).data
   # data = serializers.serialize('json', Obj.objects.filter(pk=id))
   # obj = serializers.serialize('json', obj)
   #return HttpResponse(obj, content_type="application/json")

def rdfliboutput(request, id):
   
   obj = Obj.objects.get(pk=id)
   #print(request.build_absolute_uri)
   uri = URIRef('http://127.0.0.1:8000/slg/150')
   uriobv = URIRef('http://127.0.0.1:8000/slg/150#obverse')
   urirev = URIRef('http://127.0.0.1:8000/slg/150#reverse')
   # uri = URIRef(obj.get_absolute_url)
   titel = Literal(obj.titel)
   invnr = Literal(obj.invnr)
   slgteil = Literal(obj.SlgTeil)
   stempelstellung = Literal(obj.stempelstellung)
   durchmesser = Literal(obj.durchmesser)
   nmo = Namespace("http://nomisma.org/ontology#")
   dcterms = Namespace(DCTERMS)
   
   g = Graph()
   g.bind('nmo',nmo)
   g.bind('dcterms', dcterms)
   g.add( (uri, dcterms.title, titel) )
   g.add( (uri, dcterms.identifier, invnr) )
   g.add( (uri, nmo.hasCollection, slgteil) )
   g.add( (uri, nmo.hasAxis, stempelstellung) )
   g.add( (uri, nmo.hasDiameter, durchmesser) )
   g.add( (uri, nmo.hasObverse, uriobv) )
   g.add( (uri, nmo.hasReverse, urirev) )
   # g.add( (uriobv, foaf.depiction, uriobv) )
   # g.add( (urirev, nmo.hasReverse, urirev) )

   # dcterms:identifier "2017.11.16";
	# nmo:hasCollection <http://nomisma.org/id/ans>;
	# nmo:hasAxis "12"^^<xsd:integer>;
	# nmo:hasWeight "25.59"^^xsd:decimal;
	# nmo:hasDiameter "32"^^xsd:decimal;
	# nmo:hasObverse <http://numismatics.org/collection/2017.11.16#obverse>;
	# nmo:hasReverse <http://numismatics.org/collection/2017.11.16#reverse>;
	# void:inDataset <http://numismatics.org/search/>.
   turt = g.serialize(format='turtle')
   return HttpResponse(content_type="text/turtle; charset=utf-8", content=turt)

@api_view(['POST'])
def login(request):
    user = get_object_or_404(User, username=request.data['username'])
    if not user.check_password(request.data['password']):
        return Response("missing user", status=status.HTTP_404_NOT_FOUND)
    token, created = Token.objects.get_or_create(user=user)
    serializer = UserSerializer(user)
    return Response({'token': token.key, 'user': serializer.data})

class MuenztypListCreate(generics.ListCreateAPIView):
    serializer_class = MuenztypSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = Muenztyp.objects.all()
        muenztyptitel = self.request.query_params.get('muenztyptitel', None)
        if muenztyptitel is not None:
            queryset = queryset.filter(muenztyptitel=muenztyptitel)
        link = self.request.query_params.get('link', None)
        if link is not None:
            queryset = queryset.filter(link=link)
        return queryset
        
    def perform_create(self, serializer):
        instance = serializer.save()  # Speichert das Objekt
        process_download_infos(instance)      # Führt Ihre benutzerdefinierte Funktion aus

def _clean_concordia_value(value):
    value = "" if value is None else str(value).strip()
    return value or None

def _nomisma_candidates(value):
    clean = _clean_concordia_value(value)
    if not clean:
        return []
    values = [clean]
    if not clean.startswith("http://") and not clean.startswith("https://"):
        values.append(f"http://nomisma.org/id/{clean}")
        values.append(f"https://nomisma.org/id/{clean}")
    return values

def _lookup_by_name_or_nomisma(model, value):
    clean = _clean_concordia_value(value)
    if not clean:
        return None
    candidates = _nomisma_candidates(clean)
    if hasattr(model, "name_nom_id"):
        obj = model.objects.filter(name_nom_id__in=candidates).first()
        if obj:
            return obj
    obj = model.objects.filter(name__iexact=clean).first()
    if obj:
        return obj
    if hasattr(model, "name_nom_id"):
        suffix = clean.rstrip("/").split("/")[-1]
        obj = model.objects.filter(name_nom_id__iendswith=f"/{suffix}").first()
        if obj:
            return obj
    return model.objects.filter(name__icontains=clean).first()

def _lookup_by_id_or_name(model, value, label=None):
    raw_id = None
    raw_label = label
    if isinstance(value, dict):
        raw_id = value.get("id") or value.get("target_id") or value.get("pk")
        raw_label = raw_label or value.get("name") or value.get("label") or value.get("target_label")
    else:
        raw_id = value

    try:
        if raw_id not in (None, ""):
            obj = model.objects.filter(pk=int(raw_id)).first()
            if obj:
                return obj
    except (TypeError, ValueError):
        pass

    return _lookup_by_name_or_nomisma(model, raw_label if raw_label not in (None, "") else value)

def _lookup_ref(value, ref_id=None):
    clean_id = _clean_concordia_value(ref_id)
    if clean_id:
        try:
            obj = Ref.objects.filter(pk=int(clean_id)).first()
            if obj:
                return obj
        except (TypeError, ValueError):
            pass
    clean = _clean_concordia_value(value)
    if not clean:
        return None
    obj = Ref.objects.filter(abk__iexact=clean).first()
    if obj:
        return obj
    return Ref.objects.filter(abk__icontains=clean).first()

def _safe_int(value):
    clean = _clean_concordia_value(value)
    if clean is None:
        return None
    try:
        return int(clean)
    except (TypeError, ValueError):
        return None

def _add_concordia_person(muenztyp, raw_value, funktion_id, appears_on_rev, unresolved):
    clean = _clean_concordia_value(raw_value)
    if not clean:
        return
    values = [part.strip() for part in re.split(r"[,;]", clean) if part.strip()]
    for value in values:
        person = _lookup_by_name_or_nomisma(Person, value)
        if not person:
            unresolved.append({"field": "person", "value": value})
            continue
        funktion = PersonFunktion.objects.filter(pk=funktion_id).first()
        if not funktion:
            unresolved.append({"field": "person_function", "value": funktion_id})
            continue
        Mztyp_Person.objects.get_or_create(
            Mztyp=muenztyp,
            idfk_Person=person,
            idfk_PersonFunktion=funktion,
            appears_on_rev=appears_on_rev,
        )

class ConcordiaMuenztypCreateView(APIView):
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def post(self, request):
        payload = request.data or {}
        unresolved = []
        muenztyptitel = _clean_concordia_value(payload.get("muenztyptitel") or payload.get("prefLabel"))
        if not muenztyptitel:
            return Response({"ok": False, "error": "muenztyptitel or prefLabel is required"}, status=400)

        with transaction.atomic():
            obj = Muenztyp(
                muenztyptitel=muenztyptitel,
                titel=_clean_concordia_value(payload.get("titel") or payload.get("generated_title")),
                link=_clean_concordia_value(payload.get("link") or payload.get("subject_base")),
                dat_von=_safe_int(payload.get("dat_von") or payload.get("hasStartDate")),
                dat_bis=_safe_int(payload.get("dat_bis") or payload.get("hasEndDate")),
                avleg=_clean_concordia_value(payload.get("avleg") or payload.get("legend_obv")),
                rvleg=_clean_concordia_value(payload.get("rvleg") or payload.get("legend_rev")),
                avbeschr=_clean_concordia_value(payload.get("avbeschr") or payload.get("desc_obv")),
                rvbeschr=_clean_concordia_value(payload.get("rvbeschr") or payload.get("desc_rev")),
                nummer=_clean_concordia_value(payload.get("nummer")),
            )

            obj.Nominal = _lookup_by_name_or_nomisma(Nominal, payload.get("Nominal") or payload.get("hasDenomination"))
            obj.Mzstaette = _lookup_by_name_or_nomisma(Mzstaette, payload.get("Mzstaette") or payload.get("hasMint"))
            obj.Muenzstand = _lookup_by_name_or_nomisma(Muenzstand, payload.get("Muenzstand") or payload.get("hasMuenzstand") or "Antike Herrscherprägung")
            obj.Ref = _lookup_ref(payload.get("Ref"), payload.get("RefId"))
            obj.av_bildtyp = _lookup_by_id_or_name(AvBildtyp, payload.get("av_bildtyp"), payload.get("av_bildtyp_label"))
            obj.rv_bildtyp = _lookup_by_id_or_name(RvBildtyp, payload.get("rv_bildtyp"), payload.get("rv_bildtyp_label"))
            obj.rv_beizeichen = _lookup_by_name_or_nomisma(RvBeizeichen, payload.get("rv_beizeichen") or payload.get("rv_beizeichen_suggestion") or payload.get("mintMark"))
            if "OffizinSymbol" in globals():
                obj.rv_offizin_symbol = _lookup_by_name_or_nomisma(OffizinSymbol, payload.get("rv_offizin_symbol") or payload.get("rv_offizin_symbol_suggestion") or payload.get("officinaMark"))

            for field_name, raw_value, resolved in [
                ("hasDenomination", payload.get("hasDenomination"), obj.Nominal),
                ("hasMint", payload.get("hasMint"), obj.Mzstaette),
                ("hasMuenzstand", payload.get("hasMuenzstand"), obj.Muenzstand),
                ("Ref", payload.get("Ref"), obj.Ref),
                ("av_bildtyp", payload.get("av_bildtyp") or payload.get("av_bildtyp_label") or payload.get("desc_obv"), obj.av_bildtyp),
                ("rv_bildtyp", payload.get("rv_bildtyp") or payload.get("rv_bildtyp_label") or payload.get("desc_rev"), obj.rv_bildtyp),
                ("rv_beizeichen", payload.get("rv_beizeichen_suggestion") or payload.get("mintMark"), obj.rv_beizeichen),
                ("rv_offizin_symbol", payload.get("rv_offizin_symbol_suggestion") or payload.get("officinaMark"), getattr(obj, "rv_offizin_symbol", None)),
            ]:
                if _clean_concordia_value(raw_value) and not resolved:
                    unresolved.append({"field": field_name, "value": raw_value})

            try:
                obj.save()
            except IntegrityError:
                return Response({"ok": False, "error": "Münztyp mit diesem Titel existiert bereits"}, status=409)
            _add_concordia_person(obj, payload.get("hasAuthority"), 1, False, unresolved)
            _add_concordia_person(obj, payload.get("hasIssuer"), 1, False, unresolved)
            _add_concordia_person(obj, payload.get("dargestellt_av_eligius") or payload.get("dargestellt_av"), 2, False, unresolved)
            _add_concordia_person(obj, payload.get("dargestellt_rv"), 2, True, unresolved)

        return Response({
            "ok": True,
            "created_type_id": obj.id,
            "type": MuenztypFilterSerializer([obj], many=True).data[0],
            "unresolved": unresolved,
        }, status=201)

class MuenztypFilterView(APIView):
    """Return Muenztypen in Concordia type_filter format with rich filtering."""
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get(self, request):
        qs = Muenztyp.objects.select_related(
            'Nominal', 'Mzstaette', 'av_bildtyp', 'rv_bildtyp', 'Ref',
            'av_beizeichen', 'av_offizin_symbol',
        )

        # --- text search across title, citation and represented person ---
        q = (request.query_params.get('q') or '').strip()
        if q:
            for word in q.split():
                word = word.strip()
                if word:
                    qs = qs.filter(
                        Q(muenztyptitel__icontains=word)
                        | Q(titel__icontains=word)
                        | Q(Ref__abk__icontains=word)
                        | Q(Ref__zitat__icontains=word)
                        | Q(nummer__icontains=word)
                        | Q(mztyp_person__idfk_Person__name__icontains=word)
                    ).distinct()

        # --- legend filters ---
        for param, field in [('legend_obv', 'avleg'), ('legend_rev', 'rvleg')]:
            raw = (request.query_params.get(param) or '').strip()
            if raw:
                for term in [t.strip() for t in raw.split(',') if t.strip()]:
                    pattern = self._wildcard_to_lookup(term)
                    qs = qs.filter(**{f'{field}__{pattern[0]}': pattern[1]})

        # --- description filters ---
        for param, field in [('desc_obv', 'avbeschr'), ('desc_rev', 'rvbeschr')]:
            raw = (request.query_params.get(param) or '').strip()
            if raw:
                for term in [t.strip() for t in raw.split(',') if t.strip()]:
                    qs = qs.filter(**{f'{field}__icontains': term})

        # --- dargestellt_search (person name via Mztyp_Person) ---
        dargestellt_search = (request.query_params.get('dargestellt_search') or '').strip()
        if dargestellt_search:
            qs = qs.filter(
                mztyp_person__idfk_Person__name__icontains=dargestellt_search,
            ).distinct()

        # --- denomination (Nominal.name) ---
        denominations = request.query_params.getlist('denomination')
        denominations = [d.strip() for d in denominations if d.strip()]
        if denominations:
            qs = qs.filter(Nominal__name__in=denominations)

        # --- mintMark (av_beizeichen.name) ---
        mint_marks = request.query_params.getlist('mintMark')
        mint_marks = [m.strip() for m in mint_marks if m.strip()]
        if mint_marks:
            qs = qs.filter(av_beizeichen__name__in=mint_marks)

        # --- mintMark_search (text search) ---
        mintmark_search = (request.query_params.get('mintMark_search') or '').strip()
        if mintmark_search:
            qs = qs.filter(av_beizeichen__name__iregex=self._wildcard_mintmark_regex(mintmark_search))

        # --- officinaMark (av_offizin_symbol.name) ---
        officina_marks = request.query_params.getlist('officinaMark')
        officina_marks = [o.strip() for o in officina_marks if o.strip()]
        if officina_marks:
            qs = qs.filter(av_offizin_symbol__name__in=officina_marks)

        # --- dargestellt pills (exact person name) ---
        dargestellt_values = request.query_params.getlist('dargestellt')
        dargestellt_values = [d.strip() for d in dargestellt_values if d.strip()]
        if dargestellt_values:
            qs = qs.filter(
                mztyp_person__idfk_Person__name__in=dargestellt_values,
            ).distinct()

        # --- description delta filters ---
        for param, field in [('obv_desc_delta', 'avbeschr'), ('rev_desc_delta', 'rvbeschr')]:
            for delta in request.query_params.getlist(param):
                delta = (delta or '').strip().lower()
                if len(delta) < 2:
                    continue
                sign, word = delta[0], delta[1:]
                if not word:
                    continue
                if sign == '+':
                    qs = qs.filter(**{f'{field}__icontains': word})
                elif sign == '-':
                    qs = qs.exclude(**{f'{field}__icontains': word})

        # --- stats ---
        stats = qs.aggregate(
            total=Count('id'),
            range_start=Min('dat_von'),
            range_end=Max('dat_bis'),
        )
        total = stats['total'] or 0

        # --- limit ---
        try:
            limit = max(1, int(request.query_params.get('limit', '5000')))
        except (TypeError, ValueError):
            limit = 5000

        type_ids = list(qs.order_by('muenztyptitel').values_list('id', flat=True)[:limit])
        types_qs = Muenztyp.objects.filter(id__in=type_ids).select_related(
            'Nominal', 'Mzstaette', 'av_bildtyp', 'rv_bildtyp', 'Ref',
            'av_beizeichen', 'av_offizin_symbol',
        ).order_by('muenztyptitel')

        # Prefetch person names for dargestellt_av_eligius
        person_map = {}
        mztp_rows = Mztyp_Person.objects.filter(
            Mztyp_id__in=type_ids,
        ).select_related('idfk_Person').values_list('Mztyp_id', 'idfk_Person__name')
        for mztyp_id, pname in mztp_rows:
            person_map.setdefault(mztyp_id, []).append(pname)

        results = []
        for obj in types_qs:
            obj._dargestellt_names = ', '.join(person_map.get(obj.id, []))
            results.append(obj)

        serializer = MuenztypFilterSerializer(results, many=True)
        return Response({
            'ok': True,
            'total': total,
            'returned': len(serializer.data),
            'limit': limit,
            'range_start': stats.get('range_start'),
            'range_end': stats.get('range_end'),
            'types': serializer.data,
        })

    @staticmethod
    def _wildcard_to_lookup(raw_term):
        """Convert Concordia wildcard syntax (^, €, ?) to Django ORM lookup."""
        starts_with = raw_term.startswith('^')
        ends_with = raw_term.endswith('€') and len(raw_term) > 1
        core = raw_term
        if starts_with:
            core = core[1:]
        if ends_with:
            core = core[:-1]
        core = core.replace('?', '_')

        if starts_with and ends_with:
            return ('iexact', core)
        if starts_with:
            return ('istartswith', core)
        if ends_with:
            return ('iendswith', core)
        return ('icontains', core)

    @staticmethod
    def _wildcard_mintmark_regex(raw_term):
        """Convert mintMark wildcard syntax so ? matches one char or a [symbol]."""
        starts_with = raw_term.startswith('^')
        ends_with = raw_term.endswith('€') and len(raw_term) > 1
        core = raw_term
        if starts_with:
            core = core[1:]
        if ends_with:
            core = core[:-1]

        parts = []
        for char in core:
            if char == '?':
                parts.append(r'(\[[^]]+\]|.)')
            else:
                parts.append(re.escape(char))

        pattern = ''.join(parts)
        if starts_with:
            pattern = f'^{pattern}'
        if ends_with:
            pattern = f'{pattern}$'
        return pattern


class ObjektList(generics.ListAPIView):
    serializer_class = ObjInventorySerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = Obj.objects.select_related('Typ', 'Typ__workflow', 'workflow', 'rv_offizin')
        invnr = self.request.query_params.get('invnr')
        if invnr:
            queryset = queryset.filter(invnr__iexact=invnr)
        return queryset

class ObjektDetail(generics.RetrieveUpdateAPIView):
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'invnr'
    lookup_url_kwarg = 'invnr'

    def get_queryset(self):
        return Obj.objects.select_related(
            'Typ', 'Typ__workflow', 'workflow',
            'idfk_Muenzstand', 'idfk_Nominal', 'idfk_Mzstaette',
            'av_bildtyp', 'av_beizeichen', 'av_offizin', 
            'rv_bildtyp', 'rv_beizeichen', 'rv_offizin'
        )

    def get_serializer_class(self):
        if self.request.method in ('PUT', 'PATCH'):
            return ObjDetailUpdateSerializer
        return ObjDetailSerializer

    def update(self, request, *args, **kwargs):
        # Bei PATCH immer partial=True, damit nur gesendete Felder aktualisiert werden
        partial = kwargs.pop('partial', False) or request.method == 'PATCH'
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)

        read_serializer = ObjDetailSerializer(instance, context=self.get_serializer_context())
        return Response(read_serializer.data)

class RefList(generics.ListCreateAPIView):
    serializer_class = RefSerializer

    def get_queryset(self):
        queryset = Ref.objects.all()
        abk = self.request.query_params.get('abk', None)
        if abk is not None:
            queryset = queryset.filter(abk=abk)
        return queryset
    
class AvBildtypList(generics.ListCreateAPIView):
    serializer_class = AvBildtypSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = AvBildtyp.objects.all()
        query = self.request.query_params.get('query', None)
        schlagwort = self.request.query_params.get('schlagwort', None)

        if query is not None:
            # Mehrere Suchbegriffe erlauben - Anführungszeichen respektieren
            # "geht n. l." bleibt zusammen, normale Wörter werden getrennt
            try:
                # shlex.split respektiert Anführungszeichen
                tokens = shlex.split(query.strip())
            except ValueError:
                # Falls shlex fehlschlägt, auf einfaches Splitting zurückfallen
                tokens = [t for t in re.split(r"[\s,]+", query.strip()) if t]
            
            for token in tokens:
                queryset = queryset.filter(
                    Q(name__icontains=token) |
                    Q(abk__icontains=token) |
                    Q(schlagworte__name__icontains=token)
                )

        if schlagwort is not None:
            queryset = queryset.filter(schlagworte__name__icontains=schlagwort)

        return queryset.distinct()

class AvBildtypDetail(generics.RetrieveUpdateAPIView):
    serializer_class = AvBildtypSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = AvBildtyp.objects.all()
    
class RvBildtypList(generics.ListCreateAPIView):
    serializer_class = RvBildtypSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = RvBildtyp.objects.all()
        query = self.request.query_params.get('query', None)

        if query is not None:
            # Zuerst exakte Übereinstimmung prüfen (ganzer String)
            exact_match_queryset = queryset.filter(
                Q(name__iexact=query) | Q(abk__iexact=query)
            )
            if exact_match_queryset.exists():
                return exact_match_queryset

            # Mehrere Suchbegriffe erlauben - Anführungszeichen respektieren
            # "geht n. l." bleibt zusammen, normale Wörter werden getrennt
            try:
                # shlex.split respektiert Anführungszeichen
                tokens = shlex.split(query.strip())
            except ValueError:
                # Falls shlex fehlschlägt, auf einfaches Splitting zurückfallen
                tokens = [t for t in re.split(r"[\s,]+", query.strip()) if t]
            
            for token in tokens:
                queryset = queryset.filter(
                    Q(name__icontains=token) |
                    Q(abk__icontains=token) |
                    Q(schlagworte__name__icontains=token)
                )

        return queryset.distinct()

class RvBildtypDetail(generics.RetrieveUpdateAPIView):
    serializer_class = RvBildtypSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = RvBildtyp.objects.all()

class MzstaetteList(generics.ListCreateAPIView):
    serializer_class = MzstaettenSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticatedOrReadOnly]

    def get_queryset(self):
        queryset = Mzstaette.objects.all()
        query = self.request.query_params.get('query', None)
        name_nom_id = self.request.query_params.get('name_nom_id', None)
        nomisma_id = self.request.query_params.get('nomisma_id', None)

        if query is not None:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(name_nom_id__icontains=query)
            )

        # Suche nach Nomisma ID - bevorzuge exakte Übereinstimmung
        if nomisma_id is not None:
            queryset = queryset.filter(name_nom_id__iexact=nomisma_id)
        elif name_nom_id is not None:
            queryset = queryset.filter(name_nom_id__iexact=name_nom_id)

        return queryset.distinct()

class MzstaetteDetail(generics.RetrieveUpdateAPIView):
    serializer_class = MzstaettenSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = Mzstaette.objects.all()

class NominalList(generics.ListCreateAPIView):
    serializer_class = NominalSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticatedOrReadOnly]

    def get_queryset(self):
        queryset = Nominal.objects.all()
        query = self.request.query_params.get('query', None)
        name_nom_id = self.request.query_params.get('name_nom_id', None)

        if query is not None:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(name_nom_id__icontains=query)
            )

        if name_nom_id is not None:
            queryset = queryset.filter(name_nom_id__iexact=name_nom_id)

        return queryset.distinct()

class NominalDetail(generics.RetrieveUpdateAPIView):
    serializer_class = NominalSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = Nominal.objects.all()

class HerstellungsmerkmaleList(generics.ListCreateAPIView):
    serializer_class = HerstellungsmerkmaleSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Herstellungsmerkmale.objects.all()

class Sek_MerkmaleList(generics.ListCreateAPIView):
    serializer_class = Sek_MerkmaleSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Sek_Merkmale.objects.all()

class PersonList(generics.ListCreateAPIView):
    serializer_class = PersonSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = Person.objects.all()
        query = self.request.query_params.get('query', None)
        name_nom_id = self.request.query_params.get('name_nom_id', None)

        if query is not None:
            queryset = queryset.filter(
                Q(name__icontains=query) | Q(name_nom_id__icontains=query)
            )

        if name_nom_id is not None:
            queryset = queryset.filter(name_nom_id__iexact=name_nom_id)

        return queryset.distinct()

class PersonCoinImagesView(generics.ListAPIView):
    serializer_class = PersonCoinImageSerializer

    def get_queryset(self):
        name = self.request.query_params.get('name', None)
        queryset = Person.objects.all()
        if name:
            queryset = queryset.filter(
                name__icontains=name,
                mztyp_person__idfk_PersonFunktion=2  # Hardcoded function or adjust as needed
            ).distinct()
        return queryset

class MuenzstandList(generics.ListCreateAPIView):
    serializer_class = MuenzstandSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = Muenzstand.objects.all()
        query = self.request.query_params.get('query', None)
        if query is not None:
            queryset = queryset.filter(name__icontains=query)
        return queryset.order_by('name')

class MuenzstandDetail(generics.RetrieveUpdateAPIView):
    serializer_class = MuenzstandSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = Muenzstand.objects.all()

class AvBeizeichenList(generics.ListCreateAPIView):
    serializer_class = AvBeizeichenSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = AvBeizeichen.objects.all()
        query = self.request.query_params.get('query', None)
        if query is not None:
            queryset = queryset.filter(name__icontains=query)
        return queryset.order_by('name')

class AvBeizeichenDetail(generics.RetrieveUpdateAPIView):
    serializer_class = AvBeizeichenSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = AvBeizeichen.objects.all()

class AvOffizinList(generics.ListCreateAPIView):
    serializer_class = AvOffizinSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = AvOffizin.objects.all()
        query = self.request.query_params.get('query', None)
        if query is not None:
            queryset = queryset.filter(name__icontains=query)
        return queryset.order_by('name')

class AvOffizinDetail(generics.RetrieveUpdateAPIView):
    serializer_class = AvOffizinSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = AvOffizin.objects.all()

class RvBeizeichenList(generics.ListCreateAPIView):
    serializer_class = RvBeizeichenSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = RvBeizeichen.objects.all()
        query = self.request.query_params.get('query', None)
        if query is not None:
            queryset = queryset.filter(name__icontains=query)
        return queryset.order_by('name')

class RvBeizeichenDetail(generics.RetrieveUpdateAPIView):
    serializer_class = RvBeizeichenSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = RvBeizeichen.objects.all()

class RvOffizinList(generics.ListCreateAPIView):
    serializer_class = RvOffizinSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = RvOffizin.objects.all()
        query = self.request.query_params.get('query', None)
        if query is not None:
            queryset = queryset.filter(name__icontains=query)
        return queryset.order_by('name')

class RvOffizinDetail(generics.RetrieveUpdateAPIView):
    serializer_class = RvOffizinSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = RvOffizin.objects.all()

class WorkflowList(generics.ListAPIView):
    serializer_class = WorkflowSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = Workflow.objects.all()
        query = self.request.query_params.get('query', None)
        if query is not None:
            queryset = queryset.filter(name__icontains=query)
        return queryset.order_by('reihenfolge')

class WorkflowDetail(generics.RetrieveAPIView):
    serializer_class = WorkflowSerializer
    authentication_classes = [
        QueryParamTokenAuthentication,
        SessionAuthentication,
        TokenAuthentication,
    ]
    permission_classes = [IsAuthenticated]
    lookup_field = 'pk'
    queryset = Workflow.objects.all()

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen
def collections_api(request):
    qs = Slg.objects.select_related('kategorie').all()
    return Response(SlgSerializer(qs, many=True).data)

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen
def stats_api(request):
    data = cache.get_or_set(
        'index_stats',
        lambda: {
            "anzahl_qs": Obj.objects.count(),
            "anzahl_mt": Muenztyp.objects.count(),
            "anzahl_person": Person.objects.count(),
            "anzahl_mstaetten": Mzstaette.objects.count(),
            "max_min_gewicht": Obj.objects.aggregate(gewicht__min=Min('gewicht'), gewicht__max=Max('gewicht')),
            "max_min_datierung": Muenztyp.objects.aggregate(dat_von__min=Min('dat_von'), dat_von__max=Max('dat_von')),
        },
        3600
    )
    return Response(data)

class CoinsContextAPI(APIView):
    permission_classes = [AllowAny] # Or IsAuthenticated if you prefer

    def get(self, request, pk):
        # Basic existence check for the main object
        # You might want to add more specific permission checks here
        # or reuse logic from your existing detail view if applicable.
        get_object_or_404(Obj.objects.select_related('Typ'), pk=pk)

        # The 'request' object passed to get_matrix_from_muenztypen
        # will contain URL parameters. Your apply_filters function
        # should be able to use these to filter Muenztyp.
        matrix = get_matrix_from_muenztypen(request, pk)
        return Response(matrix, status=200)

@api_view(['GET'])
@permission_classes([AllowAny])  # Diese Zeile hinzufügen (falls noch nicht vorhanden)
def cycle_coin(request):
    """
    /api/cycle-coin/?type_id=<mztyp>&current_obj_id=<obj>&direction=next|prev
    liefert {next_obj_id, card_html}
    """
    mt_id      = request.GET.get('type_id')
    current_id = request.GET.get('current_obj_id')
    direction  = request.GET.get('direction', 'next')

    if not (mt_id and current_id):
        return Response({"detail": "type_id oder current_obj_id fehlt"}, status=400)

    objs = list(
        Obj.objects
            .filter(Typ_id=mt_id)
            .only('id', 'Typ')                # schlank
            .order_by('id')
    )

    if not objs:
        return Response({"detail": "Keine Objekte für den angegebenen type_id gefunden."}, status=404)

    ids = [o.id for o in objs]

    try:
        idx = ids.index(int(current_id))
    except ValueError:                        # falls Karte nicht (mehr) in der Liste
        idx = -1

    nxt = (idx + 1) % len(objs) if direction == 'next' else (idx - 1) % len(objs)
    next_obj = objs[nxt]

    html = render_to_string(
        "slg/partials/coin_preview_card.html",
        {"obj": next_obj, "light": True}
    )
    return Response({"next_obj_id": next_obj.id, "card_html": html})

def ereignisse_api(request):
    """
    REST API-Endpunkt für Ereignisse (slginformation).

    Parameter werden über die GET-Parameter übergeben.
    """
    cache_key = f"timeline_api:{request.get_full_path()}"
    data = cache.get(cache_key)
    if data is None:
        data = get_timeline_data(request)
        cache.set(cache_key, data, 300)  # 5 Minuten
    return JsonResponse(data)


def timeline_view(request):
    """
    Rendert die Timeline-Ansicht.
    
    Die Daten werden asynchron über die REST API geladen.
    """
    # Titel für die Seite bestimmen
    title = "Ereigniszeitleiste"
    
    # Parameter für die Timeline
    context = {
        'title': title,
        'api_url': '/api/ereignisse/',
    }
    
    return render(request, 'slg/timeline.html', context)

def prepare_coin_labels_html(queryset):
    """
    Bereitet die Daten für HTML-Unterlagszettel vor
    """
    from django.template.loader import render_to_string
    from django.http import HttpResponse
    
    # Queryset deduplizieren und optimieren
    queryset = queryset.distinct().select_related(
        'Slg', 'SlgTeil', 'Typ', 'idfk_Nominal', 'Metall', 'idfk_Mzstaette', 
        'region', 'av_bildtyp', 'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen',
        'av_offizin', 'rv_offizin', 'av_offizin_symbol', 'rv_offizin_symbol',
        'Typ__Mzstaette', 'Typ__Nominal', 'Typ__Metall', 'Typ__Ref', 'Typ__Reichskreis',
        'Typ__Muenzstand', 'idfk_Muenzstand', 'Typ__av_beizeichen', 'Typ__rv_beizeichen',
        'Typ__av_offizin_symbol', 'Typ__rv_offizin_symbol', 'fund'
    )
    
    coins_data = []
    
    for obj in queryset:
        mztyp = obj.Typ
        
        # Systematische Datensammlung orientiert am Export-System
        titel = obj.titel or (mztyp.titel if mztyp else None) or ""
        
        # Nominal - vollständig ausgeschrieben (robustere Behandlung)
        nominal_obj = obj.idfk_Nominal or (mztyp.Nominal if mztyp else None)
        if nominal_obj and str(nominal_obj) != "None":
            nominal_str = str(nominal_obj)
        else:
            nominal_str = ""
        
        # Münzstätte bestimmen (robustere Behandlung)  
        muenzstaette_obj = obj.idfk_Mzstaette or (mztyp.Mzstaette if mztyp else None)
        if muenzstaette_obj and str(muenzstaette_obj) != "None":
            muenzstaette_str = str(muenzstaette_obj)
        else:
            muenzstaette_str = ""
        
        # Datierung - orientiert am Export-System
        datierung_verbale = obj.dat_verb or (mztyp.dat_verb if mztyp else None) or ""
        datierung_von = obj.dat_von or (mztyp.dat_von if mztyp else None)
        datierung_bis = obj.dat_bis or (mztyp.dat_bis if mztyp else None)
        
        # Datierung zusammensetzen falls keine verbale Datierung vorhanden
        if not datierung_verbale and (datierung_von or datierung_bis):
            if datierung_von == datierung_bis and datierung_von:
                datierung = str(datierung_von)
            elif datierung_von and datierung_bis:
                datierung = f"{datierung_von}-{datierung_bis}"
            elif datierung_von:
                datierung = f"ab {datierung_von}"
            elif datierung_bis:
                datierung = f"bis {datierung_bis}"
            else:
                datierung = ""
        else:
            datierung = datierung_verbale
        
        # Münzstätte und Datierung kombinieren
        muenzstaette_datierung = muenzstaette_str
        if muenzstaette_str and datierung:
            muenzstaette_datierung += f", {datierung}"
        elif datierung:
            muenzstaette_datierung = datierung
        
        # Kurzzitat - orientiert am Export-System
        kurzzitat = ""
        if mztyp:
            if mztyp.muenztyptitel:
                kurzzitat = mztyp.muenztyptitel
            elif mztyp.Ref:
                kurzzitat = str(mztyp.Ref.abk) if mztyp.Ref.abk else ""
                if mztyp.nummer:
                    kurzzitat += f" {mztyp.nummer}" if kurzzitat else str(mztyp.nummer)
            # Fragezeichen für unsichere Typen hinzufügen
            if kurzzitat and obj.Typ_unsicher:
                kurzzitat += "?"
        
        # Technische Daten
        durchmesser = obj.durchmesser
        gewicht = obj.gewicht
        stempelstellung = obj.stempelstellung
        
        # Formatierung der technischen Daten
        gewicht_str = f"{gewicht:.1f}" if gewicht else ""
        durchmesser_str = f"{durchmesser:.1f}" if durchmesser else ""
        stempelstellung_str = f"{stempelstellung}" if stempelstellung else ""
        
        # Beizeichen-Logik (adaptiert aus prepare_coin_data)
        rv_beizeichen_display = ""
        rv_beizeichen_obj = obj.rv_beizeichen or (mztyp.rv_beizeichen if mztyp else None)
        rv_offizin_obj = obj.rv_offizin
        rv_offizin_symbol_obj = obj.rv_offizin_symbol or (mztyp.rv_offizin_symbol if mztyp else None)
        
        if rv_beizeichen_obj:
            beizeichen_text = rv_beizeichen_obj.name
            
            # SVG-Symbole ersetzen (für HTML-Darstellung vereinfacht)
            if rv_offizin_symbol_obj:
                # Vereinfachte Darstellung ohne HTML für Labels
                if "[Offizinszeichen]" in beizeichen_text:
                    symbol_name = getattr(rv_offizin_symbol_obj, 'name', 'Symbol')
                    beizeichen_text = beizeichen_text.replace("[Offizinszeichen]", f"[{symbol_name}]")
                if "[Monogramm]" in beizeichen_text:
                    symbol_name = getattr(rv_offizin_symbol_obj, 'name', 'Symbol')
                    beizeichen_text = beizeichen_text.replace("[Monogramm]", f"[{symbol_name}]")
            
            # Monogramm-Referenzen ersetzen (vereinfacht)
            beizeichen_text = replace_monogram_references(beizeichen_text)
            # HTML-Tags entfernen für cleane Textdarstellung
            import re
            beizeichen_text = re.sub(r'<[^>]+>', '', beizeichen_text)
            
            # ? durch Offizin-Namen ersetzen
            if "?" in beizeichen_text and rv_offizin_obj:
                beizeichen_text = beizeichen_text.replace("?", rv_offizin_obj.name)
            
            rv_beizeichen_display = beizeichen_text
        elif rv_offizin_obj:
            rv_beizeichen_display = rv_offizin_obj.name
        
        # Fundnummer und Fundjahr (neu hinzugefügt)
        fund_nr_jahr = ""
        try:
            fund = obj.fund if hasattr(obj, 'fund') else None
            if fund:
                if fund.fundnummer and fund.fundjahr:
                    fund_nr_jahr = f"Fdnr./Jahr {fund.fundnummer}/{fund.fundjahr}"
                elif fund.fundnummer:
                    fund_nr_jahr = f"Fdnr. {fund.fundnummer}"
                elif fund.fundjahr:
                    fund_nr_jahr = f"Jahr {fund.fundjahr}"
        except:
            fund_nr_jahr = ""
        
        coin_data = {
            'titel': titel,
            'nominal': nominal_str,
            'muenzstaette_datierung': muenzstaette_datierung,
            'kurzzitat': kurzzitat,
            'rv_beizeichen': rv_beizeichen_display,
            'gewicht': gewicht_str,
            'durchmesser': durchmesser_str,
            'stempelstellung': stempelstellung_str,
            'inventarnummer': str(obj.invnr),
            'fund_nr_jahr': fund_nr_jahr,  # Neu hinzugefügt
        }
        
        coins_data.append(coin_data)
    
    # HTML mit Template rendern
    html_content = render_to_string('slg/coin_labels.html', {
        'coins': coins_data
    })
    
    # HTTP Response erstellen
    response = HttpResponse(html_content, content_type='text/html; charset=utf-8')
    response['Content-Disposition'] = 'inline; filename="unterlagszettel.html"'
    
    return response

@login_required
def catalog_sort_titles(request):
    """
    View zum Sortieren der Titel für die Katalog-Erstellung
    """
    obj_ids = request.GET.get('obj_ids', '')
    if not obj_ids:
        messages.error(request, 'Keine Objekte ausgewählt.')
        return redirect('admin:slg_obj_changelist')
    
    try:
        obj_id_list = [int(id_str) for id_str in obj_ids.split(',') if id_str.strip()]
    except ValueError:
        messages.error(request, 'Ungültige Objekt-IDs.')
        return redirect('admin:slg_obj_changelist')
    
    # Alle verwendeten Titel sammeln
    obj_queryset = Obj.objects.filter(id__in=obj_id_list).select_related('Typ')
    used_titles = set()
    
    for obj in obj_queryset:
        # Titel vom Objekt oder vom Typ verwenden
        titel = obj.titel or (obj.Typ.titel if obj.Typ else None)
        if titel and titel.strip():
            used_titles.add(titel.strip())
    
    # Titel in der Datenbank erstellen/aktualisieren
    from .models import Titel
    titel_objects = []
    for titel_name in used_titles:
        titel_obj, created = Titel.objects.get_or_create(
            name=titel_name,
            defaults={'sortierung': 0}
        )
        titel_objects.append(titel_obj)
    
    # Titel nach aktueller Sortierung laden
    titel_objects = Titel.objects.filter(name__in=used_titles).order_by('sortierung', 'name')
    
    if request.method == 'POST' and request.POST.get('action') == 'save_order':
        # Sortierung speichern
        titel_ids = request.POST.getlist('titel_order[]')
        for index, titel_id in enumerate(titel_ids):
            try:
                titel = Titel.objects.get(id=int(titel_id))
                titel.sortierung = index
                titel.save()
            except (ValueError, Titel.DoesNotExist):
                continue
        
        messages.success(request, 'Titel-Reihenfolge gespeichert.')
        # Weiterleitung zur Katalog-Generierung
        return redirect(f"/catalog/generate/?obj_ids={obj_ids}")
    
    context = {
        'titel_objects': titel_objects,
        'obj_ids': obj_ids,
    }
    
    return render(request, 'slg/catalog_sort_titles.html', context)

@login_required
def catalog_generate(request):
    """
    View zur Generierung des Katalogs als HTML
    """
    obj_ids = request.GET.get('obj_ids', '')
    if not obj_ids:
        messages.error(request, 'Keine Objekte ausgewählt.')
        return redirect('admin:slg_obj_changelist')
    
    try:
        obj_id_list = [int(id_str) for id_str in obj_ids.split(',') if id_str.strip()]
    except ValueError:
        messages.error(request, 'Ungültige Objekt-IDs.')
        return redirect('admin:slg_obj_changelist')
    
    # Objekte laden mit optimierten Queries
    obj_queryset = Obj.objects.filter(id__in=obj_id_list).select_related(
        'Slg', 'SlgTeil', 'Typ', 'idfk_Nominal', 'Metall', 'idfk_Mzstaette', 
        'region', 'av_bildtyp', 'rv_bildtyp', 'av_beizeichen', 'rv_beizeichen',
        'av_offizin', 'rv_offizin', 'av_offizin_symbol', 'rv_offizin_symbol',
        'Typ__Mzstaette', 'Typ__Nominal', 'Typ__Metall', 'Typ__Ref', 'Typ__Reichskreis',
        'Typ__Muenzstand', 'idfk_Muenzstand', 'Typ__av_beizeichen', 'Typ__rv_beizeichen',
        'Typ__av_offizin_symbol', 'Typ__rv_offizin_symbol', 'faelschung'
    ).prefetch_related(
        'Herstellungsmerkmale', 'sekundaere_Merkmale'
    )
    
    # Titel sortiert laden
    from .models import Titel
    titel_objects = Titel.objects.all().order_by('sortierung', 'name')
    
    # Objekte nach Titeln gruppieren (Original-Objekte behalten für Sortierung)
    from collections import defaultdict
    grouped_original_objects = defaultdict(list)
    
    for obj in obj_queryset:
        titel = obj.titel or (obj.Typ.titel if obj.Typ else None)
        if titel and titel.strip():
            grouped_original_objects[titel.strip()].append(obj)
    
    # Sortierungsfunktion für Objekte innerhalb einer Gruppe
    def sort_objects_within_group(obj):
        # Primäre Sortierung: Datierung (dat_von)
        datierung_von = obj.dat_von or (obj.Typ.dat_von if obj.Typ else None)
        datierung_bis = obj.dat_bis or (obj.Typ.dat_bis if obj.Typ else None)
        
        # Sekundäre Sortierung: Münzstätte
        muenzstaette_obj = obj.idfk_Mzstaette or (obj.Typ.Mzstaette if obj.Typ else None)
        muenzstaette_name = str(muenzstaette_obj) if muenzstaette_obj else ""
        
        # Sortierkriterien: (datierung_von, datierung_bis, münzstätte_name)
        # None-Werte werden als sehr große Zahlen behandelt, um sie ans Ende zu setzen
        return (
            datierung_von if datierung_von is not None else 9999,
            datierung_bis if datierung_bis is not None else 9999,
            muenzstaette_name
        )
    
    # Finale Gruppierung nach Titel-Sortierung mit interner Sortierung
    catalog_groups = []
    for titel_obj in titel_objects:
        if titel_obj.name in grouped_original_objects:
            # Objekte innerhalb der Gruppe sortieren
            sorted_objects = sorted(
                grouped_original_objects[titel_obj.name], 
                key=sort_objects_within_group
            )
            
            # Münzdaten für sortierte Objekte vorbereiten
            sorted_coin_data = [
                prepare_coin_data_for_catalog(obj) for obj in sorted_objects
            ]
            
            catalog_groups.append({
                'title': titel_obj.name,
                'objects': sorted_coin_data
            })
    
    # HTML rendern
    from django.template.loader import render_to_string
    html_content = render_to_string('slg/catalog.html', {
        'catalog_groups': catalog_groups
    })
    
    # HTTP Response erstellen
    response = HttpResponse(html_content, content_type='text/html; charset=utf-8')
    response['Content-Disposition'] = 'inline; filename="katalog.html"'
    
    return response

def prepare_coin_data_for_catalog(obj):
    """
    Bereitet die Münzdaten für den Katalog auf, ähnlich wie prepare_coin_data aus der Detail-Ansicht
    """
    mztyp = obj.Typ
    
    # Grunddaten
    data = {
        'invnr': obj.invnr,
        'durchmesser': obj.durchmesser,
        'gewicht': obj.gewicht,
        'stempelstellung': obj.stempelstellung,
        'abnutzung': obj.abnutzung,
        'titel': obj.titel or (mztyp.titel if mztyp else None),
        'has_typ': mztyp is not None,
        'anmerkung': obj.anmerkung or "",
    }
    
    # Avers/Revers Informationen für Objekte ohne Typ
    if not mztyp:
        data['avleg'] = obj.avleg or ""
        data['av_bildtyp'] = str(obj.av_bildtyp) if obj.av_bildtyp else ""
        data['rvleg'] = obj.rvleg or ""
        data['rv_bildtyp'] = str(obj.rv_bildtyp) if obj.rv_bildtyp else ""
    else:
        data['avleg'] = ""
        data['av_bildtyp'] = ""
        data['rvleg'] = ""
        data['rv_bildtyp'] = ""
    
    # Nominal
    nominal_obj = obj.idfk_Nominal or (mztyp.Nominal if mztyp else None)
    data['nominal'] = str(nominal_obj) if nominal_obj else ""
    
    # Münzstätte
    muenzstaette_obj = obj.idfk_Mzstaette or (mztyp.Mzstaette if mztyp else None)
    data['muenzstaette'] = str(muenzstaette_obj) if muenzstaette_obj else ""
    
    # Datierung
    datierung_verbale = obj.dat_verb or (mztyp.dat_verb if mztyp else None) or ""
    datierung_von = obj.dat_von or (mztyp.dat_von if mztyp else None)
    datierung_bis = obj.dat_bis or (mztyp.dat_bis if mztyp else None)
    
    if not datierung_verbale and (datierung_von or datierung_bis):
        if datierung_von == datierung_bis and datierung_von:
            data['datierung'] = str(datierung_von)
        elif datierung_von and datierung_bis:
            data['datierung'] = f"{datierung_von}-{datierung_bis}"
        elif datierung_von:
            data['datierung'] = f"ab {datierung_von}"
        elif datierung_bis:
            data['datierung'] = f"bis {datierung_bis}"
        else:
            data['datierung'] = ""
    else:
        data['datierung'] = datierung_verbale
    
    # Kurzzitat (Literatur)
    if mztyp:
        if mztyp.muenztyptitel:
            kurzzitat = mztyp.muenztyptitel
        elif mztyp.Ref:
            kurzzitat = str(mztyp.Ref.abk) if mztyp.Ref.abk else ""
            if mztyp.nummer:
                kurzzitat += f" {mztyp.nummer}" if kurzzitat else str(mztyp.nummer)
        else:
            kurzzitat = ""
        
        # Fragezeichen für unsichere Typen hinzufügen
        if kurzzitat and obj.Typ_unsicher:
            kurzzitat += "?"
        data['literatur'] = kurzzitat
    else:
        data['literatur'] = ""
    
    # Beizeichen-Logik für Avers (adaptiert aus prepare_coin_labels_html)
    av_beizeichen_display = ""
    av_beizeichen_obj = obj.av_beizeichen or (mztyp.av_beizeichen if mztyp else None)
    av_offizin_obj = obj.av_offizin
    av_offizin_symbol_obj = obj.av_offizin_symbol or (mztyp.av_offizin_symbol if mztyp else None)
    
    if av_beizeichen_obj:
        beizeichen_text = av_beizeichen_obj.name
        
        # SVG-Symbole ersetzen (für HTML-Darstellung vereinfacht)
        if av_offizin_symbol_obj:
            # Vereinfachte Darstellung ohne HTML für Katalog
            if "[Offizinszeichen]" in beizeichen_text:
                symbol_name = getattr(av_offizin_symbol_obj, 'name', 'Symbol')
                beizeichen_text = beizeichen_text.replace("[Offizinszeichen]", f"[{symbol_name}]")
            if "[Monogramm]" in beizeichen_text:
                symbol_name = getattr(av_offizin_symbol_obj, 'name', 'Symbol')
                beizeichen_text = beizeichen_text.replace("[Monogramm]", f"[{symbol_name}]")
        
        # Monogramm-Referenzen ersetzen (vereinfacht)
        beizeichen_text = replace_monogram_references(beizeichen_text)
        # HTML-Tags entfernen für cleane Textdarstellung
        import re
        beizeichen_text = re.sub(r'<[^>]+>', '', beizeichen_text)
        
        # ? durch Offizin-Namen ersetzen
        if "?" in beizeichen_text and av_offizin_obj:
            beizeichen_text = beizeichen_text.replace("?", av_offizin_obj.name)
        
        av_beizeichen_display = beizeichen_text
    elif av_offizin_obj:
        av_beizeichen_display = av_offizin_obj.name
    
    data['av_beizeichen'] = av_beizeichen_display
    
    # Beizeichen-Logik für Revers (adaptiert aus prepare_coin_labels_html)
    rv_beizeichen_display = ""
    rv_beizeichen_obj = obj.rv_beizeichen or (mztyp.rv_beizeichen if mztyp else None)
    rv_offizin_obj = obj.rv_offizin
    rv_offizin_symbol_obj = obj.rv_offizin_symbol or (mztyp.rv_offizin_symbol if mztyp else None)
    
    if rv_beizeichen_obj:
        beizeichen_text = rv_beizeichen_obj.name
        
        # SVG-Symbole ersetzen (für HTML-Darstellung vereinfacht)
        if rv_offizin_symbol_obj:
            # Vereinfachte Darstellung ohne HTML für Katalog
            if "[Offizinszeichen]" in beizeichen_text:
                symbol_name = getattr(rv_offizin_symbol_obj, 'name', 'Symbol')
                beizeichen_text = beizeichen_text.replace("[Offizinszeichen]", f"[{symbol_name}]")
            if "[Monogramm]" in beizeichen_text:
                symbol_name = getattr(rv_offizin_symbol_obj, 'name', 'Symbol')
                beizeichen_text = beizeichen_text.replace("[Monogramm]", f"[{symbol_name}]")
        
        # Monogramm-Referenzen ersetzen (vereinfacht)
        beizeichen_text = replace_monogram_references(beizeichen_text)
        # HTML-Tags entfernen für cleane Textdarstellung
        import re
        beizeichen_text = re.sub(r'<[^>]+>', '', beizeichen_text)
        
        # ? durch Offizin-Namen ersetzen
        if "?" in beizeichen_text and rv_offizin_obj:
            beizeichen_text = beizeichen_text.replace("?", rv_offizin_obj.name)
        
        rv_beizeichen_display = beizeichen_text
    elif rv_offizin_obj:
        rv_beizeichen_display = rv_offizin_obj.name
    
    data['rv_beizeichen'] = rv_beizeichen_display
    
    # Herstellungsmerkmale, sekundäre Merkmale und Fälschung
    herstellungsmerkmale = list(obj.Herstellungsmerkmale.all()) if hasattr(obj, 'Herstellungsmerkmale') else []
    data['herstellungsmerkmale'] = ', '.join([str(merkmal) for merkmal in herstellungsmerkmale]) if herstellungsmerkmale else ""
    
    sekundaere_merkmale = list(obj.sekundaere_Merkmale.all()) if hasattr(obj, 'sekundaere_Merkmale') else []
    data['sekundaere_merkmale'] = ', '.join([str(merkmal) for merkmal in sekundaere_merkmale]) if sekundaere_merkmale else ""
    
    data['faelschung'] = str(obj.faelschung) if obj.faelschung else ""
    
    # Fundkontext (falls vorhanden)
    try:
        fund = obj.fund if hasattr(obj, 'fund') else None
        if fund:
            fund_context_parts = []
            if fund.fundort:
                fund_context_parts.append(str(fund.fundort))
            if fund.parzelle:
                fund_context_parts.append(f"Parzelle {fund.parzelle}")
            
            # Fundnummer und Fundjahr mit "/" verbinden
            fundnr_jahr = ""
            if fund.fundnummer and fund.fundjahr:
                fundnr_jahr = f"{fund.fundnummer}/{fund.fundjahr}"
            elif fund.fundnummer:
                fundnr_jahr = str(fund.fundnummer)
            elif fund.fundjahr:
                fundnr_jahr = str(fund.fundjahr)
            
            if fundnr_jahr:
                fund_context_parts.append(fundnr_jahr)
                
            data['fundkontext'] = ', '.join(fund_context_parts)
        else:
            data['fundkontext'] = ""
    except:
        data['fundkontext'] = ""
    
    return data


def objekt_list_view_mtoa(request):
    """
    Browse-View auf Basis der MuenztypObjektAnzeige (MTOA).
    Nutzt die flache MTOA-Tabelle und MtoaPerson-Bridge für effiziente Abfragen.
    """
    invnr = request.GET.get('invnr')
    if invnr:
        obj = Obj.objects.only('id').filter(invnr=invnr).first()
        if obj:
            return redirect('Objekt', id=obj.id)
    qs, is_unbestimmt, needs_distinct = _get_filtered_mtoa_queryset(request)

    # Database-agnostic grouping for maps
    mints_qs = (
        qs.exclude(mzstaette_fk__isnull=True)
        .values(
            'mzstaette_fk__id',
            'mzstaette',
            'region',
            'mzstaette_fk__lat',
            'mzstaette_fk__long',
        )
        .annotate(
            mint_id     = F('mzstaette_fk__id'),
            mint_name   = F('mzstaette'),
            mint_region = F('region'),
            mint_lat    = F('mzstaette_fk__lat'),
            mint_lon    = F('mzstaette_fk__long'),
            obj_count   = Count('pk', distinct=needs_distinct),
        )
        .order_by('mzstaette_fk__id')
    )

    mints_for_js = [
        {
            'id': m['mint_id'],
            'name': m['mint_name'],
            'region': m['mint_region'],
            'longitude': float(m['mint_lon']) if m['mint_lon'] is not None else None,
            'latitude': float(m['mint_lat']) if m['mint_lat'] is not None else None,
            'count': m['obj_count'],
        }
        for m in mints_qs
        if m['mint_lon'] and m['mint_lat']
    ]

    page = request.GET.get('page', 1)
    paginator = Paginator(qs, 30)
    anzahl = paginator.count
    variables = request.GET.copy()
    if 'page' in variables:
        del variables['page']

    try:
        qs = paginator.page(page)
    except PageNotAnInteger:
        qs = paginator.page(1)
    except EmptyPage:
        qs = paginator.page(paginator.num_pages)

    context_key = 'unbestimmt' if is_unbestimmt else 'default'
    parameters = FILTER_PARAMETERS.get(context_key, {})

    context = {
        'Objekte': qs,
        'mints_map': json.dumps(mints_for_js),
        'Anzahl': anzahl,
        'getvars': '&{0}'.format(variables.urlencode()),
        'filter_config': parameters,
    }

    return render(request, 'slg/Objektliste_mtoa.html', context)

@login_required
def excel_ocre_enrich(request):
    import io
    import openpyxl
    from django.http import HttpResponse

    if request.method == 'POST':
        if 'excel_file' not in request.FILES or not request.POST.get('ocre_column'):
            messages.error(request, 'Bitte eine Datei hochladen und den Spaltennamen für die OCRE-Links angeben.')
            return redirect('excel_ocre_enrich')

        excel_file = request.FILES['excel_file']
        ocre_column_name = request.POST.get('ocre_column').strip()

        try:
            wb = openpyxl.load_workbook(excel_file)
            sheet = wb.active

            # Find the OCRE column index (case-insensitive and ignoring extra spaces)
            header_row = list(sheet.rows)[0]
            ocre_col_idx = None
            for idx, cell in enumerate(header_row):
                if cell.value and str(cell.value).strip().lower() == ocre_column_name.lower():
                    ocre_col_idx = idx
                    break

            if ocre_col_idx is None:
                messages.error(request, f'Spalte "{ocre_column_name}" nicht gefunden.')
                return redirect('excel_ocre_enrich')

            # Fetch all MuenzTyp data into a dictionary for fast lookup to avoid N queries
            all_types = Muenztyp.objects.exclude(link__isnull=True).exclude(link="").select_related(
                'av_bildtyp', 'rv_bildtyp', 'workflow'
            ).only(
                'link', 'avbeschr', 'rvbeschr', 'av_bildtyp__name', 'rv_bildtyp__name', 'workflow__name', 'muenztyptitel'
            )
            
            id_to_data = {}
            for t in all_types:
                if t.link and '/id/' in t.link:
                    extracted_id = t.link.split('/id/')[-1].strip('/').lower()
                    
                    # Get AV and RV Descriptions from AvBildtyp/RvBildtyp (fallback to avbeschr/rvbeschr if necessary)
                    av_text = t.av_bildtyp.name if t.av_bildtyp else (t.avbeschr or "")
                    rv_text = t.rv_bildtyp.name if t.rv_bildtyp else (t.rvbeschr or "")
                    
                    id_to_data[extracted_id] = {
                        'av': av_text,
                        'rv': rv_text,
                        'wf': t.workflow.name if t.workflow else "",
                        'titel': t.muenztyptitel or ""
                    }

            # Add new headers
            max_col = sheet.max_column
            col_avbeschr = max_col + 1
            col_rvbeschr = max_col + 2
            col_workflow = max_col + 3
            col_titel = max_col + 4

            sheet.cell(row=1, column=col_avbeschr, value="Av-Beschreibung")
            sheet.cell(row=1, column=col_rvbeschr, value="Rv-Beschreibung")
            sheet.cell(row=1, column=col_workflow, value="Workflow")
            sheet.cell(row=1, column=col_titel, value="Münztyptitel")

            # Process rows using the cached map
            for row_idx in range(2, sheet.max_row + 1):
                ocre_link_cell = sheet.cell(row=row_idx, column=ocre_col_idx + 1)
                link_val = ocre_link_cell.value

                if link_val:
                    # Extract the ID from the link, assuming format like http://numismatics.org/ocre/id/ric.1(2).aug.1A
                    link_str = str(link_val).strip()
                    if '/id/' in link_str:
                        extracted_id = link_str.split('/id/')[-1].strip('/').lower()
                        
                        data = id_to_data.get(extracted_id)
                        
                        if data:
                            sheet.cell(row=row_idx, column=col_avbeschr, value=data['av'])
                            sheet.cell(row=row_idx, column=col_rvbeschr, value=data['rv'])
                            sheet.cell(row=row_idx, column=col_workflow, value=data['wf'])
                            sheet.cell(row=row_idx, column=col_titel, value=data['titel'])

            # Save to memory and return
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            response = HttpResponse(
                output.read(), 
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            original_name = excel_file.name
            new_name = original_name.replace('.xlsx', '_angereichert.xlsx') if '.xlsx' in original_name else f'{original_name}_angereichert.xlsx'
            response['Content-Disposition'] = f'attachment; filename="{new_name}"'
            return response

        except Exception as e:
            logger.error(f"Fehler bei Excel-Verarbeitung: {e}")
            messages.error(request, f'Es ist ein Fehler aufgetreten: {str(e)}')
            return redirect('excel_ocre_enrich')

    return render(request, 'slg/excel_ocre_enrich.html')
